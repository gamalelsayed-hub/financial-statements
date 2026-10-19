"""Export the statements, ratios and classification to an Excel workbook."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .categories import GROUP_NAMES, BY_KEY
from .i18n import t, lang, ratio_text
from .reports import summary

NAVY = "12355B"
NUM = '#,##0;(#,##0);"-"'
HDR_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="F2F5F9")
TOT_FILL = PatternFill("solid", fgColor="E6EDF6")
STATUS_COLORS = {"good": "067647", "warn": "B54708", "bad": "B42318", "info": "5B6475", "na": "5B6475"}


def _gname(g):
    return GROUP_NAMES[g][0 if lang() == "ar" else 1]


class Sheet:
    def __init__(self, wb, title, meta, sub, columns, widths):
        self.ws = wb.create_sheet(title[:31])
        self.ws.sheet_view.rightToLeft = lang() == "ar"
        self.ncols = len(columns)
        self.r = 1
        for txt, size in ((meta["company"], 14), (title, 12), (sub, 10)):
            c = self.ws.cell(self.r, 1, txt)
            c.font = Font(bold=size > 10, size=size, color=NAVY)
            self.ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=self.ncols)
            c.alignment = Alignment(horizontal="center")
            self.r += 1
        self.r += 1
        for i, h in enumerate(columns, start=1):
            c = self.ws.cell(self.r, i, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HDR_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.r += 1
        for i, w in enumerate(widths, start=1):
            self.ws.column_dimensions[chr(64 + i)].width = w
        self.ws.freeze_panes = self.ws.cell(self.r, 1)

    def row(self, label, values=(), kind="item", note=None, fmt=NUM):
        ws = self.ws
        c = ws.cell(self.r, 1, ("    " + label) if kind == "item" else label)
        col = 2
        if note is not None:
            ws.cell(self.r, 2, note).alignment = Alignment(horizontal="center")
            col = 3
        for v in values:
            cell = ws.cell(self.r, col, v)
            if isinstance(v, (int, float)):
                cell.number_format = fmt
            col += 1
        if kind in ("hdr", "sub", "tot"):
            for i in range(1, self.ncols + 1):
                cell = ws.cell(self.r, i)
                cell.font = Font(bold=True, color=NAVY if kind == "hdr" else "000000")
                if kind == "sub":
                    cell.fill = SUB_FILL
                    cell.border = Border(top=Side("thin"))
                if kind == "tot":
                    cell.fill = TOT_FILL
                    cell.border = Border(top=Side("medium"), bottom=Side("double"))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        self.r += 1


def export(st, meta, accounts, path):
    wb = Workbook()
    wb.remove(wb.active)
    notes = {}

    def note(line):
        if not line.accounts:
            return ""
        if line.key not in notes:
            notes[line.key] = (len(notes) + 1, line.accounts)
        return notes[line.key][0]

    # Balance sheet
    bs = st.bs
    s = Sheet(wb, t("statement_of_fp"), meta, t("as_at", d=meta["period"]),
              ["", t("note"), t("current"), t("prior")], [48, 8, 18, 18])
    s.row(t("assets"), kind="hdr")
    for grp, tk in (("CA", "total_CA"), ("NCA", "total_NCA")):
        s.row(_gname(grp), kind="hdr")
        for l in bs[grp]:
            s.row(t(l.key), [l.amount, l.prior], note=note(l))
        s.row(t(tk), list(bs[grp + "_total"]), "sub", note="")
    s.row(t("total_assets"), list(bs["assets"]), "tot", note="")
    s.row(t("liab_equity"), kind="hdr")
    for grp, tk in (("CL", "total_CL"), ("NCL", "total_NCL")):
        if not bs[grp]:
            continue
        s.row(_gname(grp), kind="hdr")
        for l in bs[grp]:
            s.row(t(l.key), [l.amount, l.prior], note=note(l))
        s.row(t(tk), list(bs[grp + "_total"]), "sub", note="")
    s.row(t("total_liabilities"), list(bs["liabilities"]), "sub", note="")
    s.row(_gname("EQ"), kind="hdr")
    for l in bs["EQ"]:
        s.row(t(l.key), [l.amount, l.prior], note=note(l))
    s.row(t("total_EQ"), list(bs["EQ_total"]), "sub", note="")
    s.row(t("total_le"), list(bs["liabilities_equity"]), "tot", note="")

    # Income statement
    s = Sheet(wb, t("statement_of_pl"), meta, t("for_period", d=meta["period"]),
              ["", t("note"), t("amount")], [48, 8, 18])
    kinds = {"item": "item", "subtotal": "sub", "total": "tot"}
    for l in st.is_["lines"]:
        if l.kind == "item" and not l.accounts and abs(l.amount) < 1e-9:
            continue
        s.row(t(l.key), [l.amount], kinds[l.kind], note=note(l))

    # Cash flow
    cf = st.cf
    s = Sheet(wb, t("tab_cf"), meta, t("for_period", d=meta["period"]), ["", t("amount")], [60, 18])
    s.row(t("cfo_title"), kind="hdr")
    s.row(t("net_income"), [cf["net_income"]])
    if cf["op_adj"]:
        s.row(t("adjustments"), kind="item")
        for l in cf["op_adj"]:
            s.row(t(l.key), [l.amount])
    if cf["op_wc"]:
        s.row(t("wc_changes"), kind="item")
        for l in cf["op_wc"]:
            s.row(t(l.key), [l.amount])
    s.row(t("cfo"), [cf["cfo"]], "sub")
    s.row(t("cfi_title"), kind="hdr")
    for l in cf["inv"]:
        s.row(t(l.key), [l.amount])
    s.row(t("cfi"), [cf["cfi"]], "sub")
    s.row(t("cff_title"), kind="hdr")
    for l in cf["fin"]:
        s.row(t(l.key), [l.amount])
    s.row(t("cff"), [cf["cff"]], "sub")
    s.row(t("net_change"), [cf["net_change"]], "sub")
    s.row(t("cash_open"), [cf["cash_open"]])
    s.row(t("cash_close"), [cf["cash_close_calc"]], "tot")

    # Changes in equity
    eq = st.equity
    cols = [k for k in eq["cols"] if any(abs(v) > 1e-9 for v in eq["rows"][k].values())] or eq["cols"][:3]
    s = Sheet(wb, t("tab_eq"), meta, t("for_period", d=meta["period"]),
              [""] + [t(k) for k in cols] + [t("total")], [40] + [18] * (len(cols) + 1))
    for part, label, kind in (("open", "eq_open", "sub"), ("ni", "eq_ni", "item"),
                              ("other", "eq_other", "item"), ("close", "eq_close", "tot")):
        vals = [eq["rows"][k][part] for k in cols]
        s.row(t(label), vals + [sum(vals)], kind)

    # Notes
    s = Sheet(wb, t("tab_notes"), meta, t("as_at", d=meta["period"]),
              [t("col_code"), t("col_name"), t("opening"), t("movement"), t("closing")], [14, 44, 16, 16, 16])
    for key, (n, accs) in sorted(notes.items(), key=lambda kv: kv[1][0]):
        s.row(f"({n}) {t(key)}", kind="hdr")
        for a in accs:
            s.ws.cell(s.r, 1, a.code)
            s.ws.cell(s.r, 2, a.name)
            for i, v in enumerate((a.open_net, a.mov_net, a.close_net), start=3):
                s.ws.cell(s.r, i, v).number_format = NUM
            s.r += 1
        s.row(t("total"), ["", sum(a.open_net for a in accs), sum(a.mov_net for a in accs),
                           sum(a.close_net for a in accs)], "sub")

    # Ratios
    s = Sheet(wb, t("ratios_title"), meta, t("for_period", d=meta["period"]),
              [t("ratio"), t("value"), t("formula"), t("assessment"), t("meaning")], [34, 16, 46, 14, 60])
    s.row(t("summary"), kind="hdr")
    for line in summary(st):
        s.ws.cell(s.r, 1, "• " + line)
        s.ws.merge_cells(start_row=s.r, start_column=1, end_row=s.r, end_column=5)
        s.r += 1
    group = None
    for r in st.ratios:
        if r["group"] != group:
            group = r["group"]
            s.row(t("grp_" + group), kind="hdr")
        v = r["value"]
        fmt = {"%": "0.0%", "x": '0.00"x"', "days": '0" ' + t("days_unit") + '"', "amount": NUM}[r["fmt"]]
        ws = s.ws
        ws.cell(s.r, 1, ratio_text(r["key"], 0))
        c = ws.cell(s.r, 2, v if v is not None else t("st_na"))
        if v is not None:
            c.number_format = fmt
        ws.cell(s.r, 3, ratio_text(r["key"], 1)).alignment = Alignment(wrap_text=True)
        st_cell = ws.cell(s.r, 4, t("st_" + r["status"]))
        st_cell.font = Font(bold=True, color=STATUS_COLORS[r["status"]])
        ws.cell(s.r, 5, ratio_text(r["key"], 2)).alignment = Alignment(wrap_text=True)
        s.r += 1

    # Classification (audit trail)
    s = Sheet(wb, t("sheet_classification"), meta, meta["period"],
              [t("col_code"), t("col_name"), t("group"), t("col_category"), t("opening"), t("movement"), t("closing")],
              [14, 40, 26, 40, 16, 16, 16])
    for a in accounts:
        cat = BY_KEY.get(a.category)
        vals = [a.code, a.name, _gname(cat.group) if cat else "", t(a.category) if cat else "",
                a.open_net, a.mov_net, a.close_net]
        for i, v in enumerate(vals, start=1):
            c = s.ws.cell(s.r, i, v)
            if i >= 5:
                c.number_format = NUM
        s.r += 1

    wb.save(path)
