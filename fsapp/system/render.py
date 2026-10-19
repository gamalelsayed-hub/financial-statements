"""Render reports to HTML (screen and PDF) and to Excel sheets."""
import html

from .reports import Para, Table

CSS = """
body { font-family: 'Segoe UI', 'Tahoma', 'Arial'; font-size: 10pt; color: #1d2433; }
table.st { border-collapse: collapse; margin-top: 6px; }
table.st td, table.st th { padding: 3px 7px; }
table.st th { background: #12355b; color: white; font-weight: 600; }
td.num { text-align: right; white-space: nowrap; }
tr.hdr td { font-weight: 700; color: #12355b; padding-top: 8px; border-bottom: 1px solid #c9d3e0; }
tr.sub td { font-weight: 700; border-top: 1px solid #8a96a8; background: #f2f5f9; }
tr.tot td { font-weight: 800; border-top: 2px solid #12355b; border-bottom: 3px double #12355b; background: #e6edf6; }
tr.note td { color: #5b6475; font-style: italic; }
.sub { color: #5b6475; font-size: 9pt; }
.good { color: #067647; font-weight: 700; }
.warn { color: #b54708; font-weight: 700; }
.bad { color: #b42318; font-weight: 700; }
h3 { color: #12355b; margin: 10px 0 2px 0; }
"""


def fmt_value(v, f):
    if v is None:
        return ""
    if isinstance(v, str) or f == "t":
        return html.escape(str(v))
    if f == "%":
        return f"{v * 100:,.1f}%"
    if f == "x":
        return f"{v:,.2f}"
    if f == "d":
        return f"{v:,.0f}"
    if abs(v) < 0.5:
        return "-"
    s = f"{abs(v):,.0f}"
    return f"({s})" if v < 0 else s


def value_format(table, row, i):
    """Format of value i of a row: the column format, unless the row itself is a ratio/days row."""
    cf = table.col_fmts[i] if table.col_fmts and i < len(table.col_fmts) else None
    if row.fmt in ("%", "x", "d") and cf in (None, "n"):
        return row.fmt
    return cf or row.fmt


def cell_text(v, f):
    if v is None:
        return ""
    if isinstance(v, str) or f == "t":
        return str(v)
    return fmt_value(v, f)


def report_html(rep, company, wrap=True):
    out = [f'<p align="center"><b style="font-size:14pt;color:#12355b">{html.escape(company)}</b><br>'
           f'<b style="font-size:12pt;color:#12355b">{html.escape(rep.title)}</b><br>'
           f'<span class="sub">{html.escape(rep.subtitle)}</span></p>']
    for b in rep.blocks:
        if isinstance(b, Para):
            out.append(f'<p class="{b.cls}">{html.escape(b.text)}</p>')
            continue
        if b.title:
            out.append(f"<h3>{html.escape(b.title)}</h3>")
        heads = "".join(f"<th>{html.escape(str(h))}</th>" for h in b.columns)
        rows = []
        ncols = len(b.columns)
        for r in b.rows:
            if r.kind == "hdr" and not r.values:
                rows.append(f'<tr class="hdr"><td nowrap colspan="{ncols}">{html.escape(r.label)}</td></tr>')
                continue
            cells = []
            for i, v in enumerate(r.values):
                f = value_format(b, r, i)
                cls = "num" if not isinstance(v, str) and f != "t" else ""
                st = r.styles[i] if r.styles and i < len(r.styles) and r.styles[i] else ""
                if not st and isinstance(v, str) and (v.startswith("▲") or v.startswith("▼")):
                    st = "warn"
                cells.append(f'<td nowrap class="{cls} {st}">{fmt_value(v, f)}</td>')
            cells += ["<td></td>"] * (ncols - 1 - len(r.values))
            rows.append(f'<tr class="{r.kind}"><td nowrap>{html.escape(r.label)}</td>{"".join(cells)}</tr>')
        out.append(f'<table class="st" dir="rtl" width="100%" cellspacing="0"><tr>{heads}</tr>{"".join(rows)}</table>')
    body = "".join(out)
    if not wrap:
        return body
    return f'<html dir="rtl"><head><meta charset="utf-8"><style>{CSS}</style></head><body dir="rtl">{body}</body></html>'


def reports_html(reps, company):
    br = '<div style="page-break-before: always;"></div>'
    body = br.join(report_html(r, company, wrap=False) for r in reps)
    return f'<html dir="rtl"><head><meta charset="utf-8"><style>{CSS}</style></head><body dir="rtl">{body}</body></html>'


# ---------------------------------------------------------------- Excel
def write_excel(reps, company, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    navy = "12355B"
    fmts = {"n": '#,##0;(#,##0);"-"', "%": "0.0%", "x": "0.00", "d": "#,##0", "t": "@"}
    fills = {"sub": PatternFill("solid", fgColor="F2F5F9"), "tot": PatternFill("solid", fgColor="E6EDF6")}
    colors = {"good": "067647", "warn": "B54708", "bad": "B42318"}
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for rep in reps:
        name = rep.title.replace("/", "-").replace(":", " ")[:31]
        base, k = name, 2
        while name in used:
            name = f"{base[:28]} {k}"
            k += 1
        used.add(name)
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        width = max([len(b.columns) for b in rep.blocks if isinstance(b, Table)] or [2])
        r = 1
        for text, size in ((company, 14), (rep.title, 12), (rep.subtitle, 10)):
            c = ws.cell(r, 1, text)
            c.font = Font(bold=size > 10, size=size, color=navy)
            r += 1
        r += 1
        for b in rep.blocks:
            if isinstance(b, Para):
                c = ws.cell(r, 1, b.text)
                c.font = Font(italic=True, color=colors.get(b.cls, "5B6475"))
                r += 2
                continue
            if b.title:
                ws.cell(r, 1, b.title).font = Font(bold=True, size=11, color=navy)
                r += 1
            for i, h in enumerate(b.columns, start=1):
                c = ws.cell(r, i, h)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=navy)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            r += 1
            for row in b.rows:
                ws.cell(r, 1, row.label)
                for i, v in enumerate(row.values):
                    f = value_format(b, row, i)
                    c = ws.cell(r, i + 2, v)
                    if isinstance(v, (int, float)):
                        c.number_format = fmts.get(f, fmts["n"])
                    st = row.styles[i] if row.styles and i < len(row.styles) else None
                    if st in colors:
                        c.font = Font(bold=True, color=colors[st])
                if row.kind in ("hdr", "sub", "tot"):
                    for i in range(1, len(b.columns) + 1):
                        c = ws.cell(r, i)
                        c.font = Font(bold=True, color=navy if row.kind == "hdr" else "000000")
                        if row.kind in fills:
                            c.fill = fills[row.kind]
                        if row.kind == "tot":
                            c.border = Border(top=Side("medium"), bottom=Side("double"))
                r += 1
            r += 1
        ws.column_dimensions["A"].width = 44
        for i in range(2, width + 1):
            ws.column_dimensions[get_column_letter(i)].width = 15
        ws.freeze_panes = "B5"
    wb.save(path)
