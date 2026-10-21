"""Draft financial statements → Excel (editable, totals as formulas) and PDF (the company's printed layout)."""
import html

from . import draft as DR

TITLES = {"fp": "قائمة المركز المالي كما في {end}", "is": "قائمة الدخل الشامل من {start} إلى {end}",
          "cf": "قائمة التدفقات النقدية من {start} إلى {end}"}
SHEETS = {"fp": "المركز المالي", "is": "قائمة الدخل", "cf": "التدفقات النقدية"}
NUM = '#,##0;(#,##0);"-"'


def period_labels(dr, ctx, key="fp"):
    """Column dates; the income statement and cash flows have a comparative only when the previous year exists."""
    first = ctx.ytd_months(dr.end)[0]
    y, m = map(int, first.split("-"))
    has = {"fp": dr.comp is not None, "is": bool(dr.comp and dr.comp.is_ is not None), "cf": dr.cf_comp is not None,
           "notes": dr.comp is not None}[key]
    return {"start": f"1/{m}/{y}", "end": DR.date_label(dr.end),
            "cur": DR.date_label(dr.end), "comp": DR.date_label(dr.comp.end) if has else ""}


def money(v):
    if v is None:
        return ""
    if abs(v) < 0.5:
        return "-"
    s = f"{abs(v):,.0f}"
    return f"‎({s})‎" if v < 0 else s


# ------------------------------------------------------------------ Excel
def to_excel(dr, ctx, path, opts):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side

    labels = period_labels(dr, ctx)
    blocks, notes = DR.statements(dr, ctx)
    wb = openpyxl.Workbook()
    thin, double = Side(style="thin"), Side(style="double")
    bold = Font(bold=True)

    def header(ws, title):
        ws.sheet_view.rightToLeft = True
        ws["A1"] = dr.company
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = opts.get("legal_form", "")
        ws["A4"] = title
        ws["A4"].font = Font(bold=True, size=12)
        ws["A5"] = opts.get("currency", "ريال سعودي")
        for col, w in zip("ABCD", (52, 8, 17, 17)):
            ws.column_dimensions[col].width = w

    for i, (key, rows) in enumerate(blocks):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = SHEETS[key]
        lab = period_labels(dr, ctx, key)
        header(ws, TITLES[key].format(**lab))
        ws.append([])
        ws.append(["البيان", "إيضاح", lab["comp"], lab["cur"]])
        for c in ws[ws.max_row]:
            c.font = bold
            c.border = Border(bottom=thin)
            c.alignment = Alignment(horizontal="center")
        refs = {}          # label → row number (for the formulas of the totals)
        since = []         # item rows since the last heading
        for label, note, cur, comp, kind in rows:
            if kind == "blank":
                ws.append([])
                continue
            ws.append([label, note or None, comp, cur])
            r = ws.max_row
            refs[label] = r
            if kind == "hdr":
                ws.cell(r, 1).font = Font(bold=True, underline="single")
                since = []
                continue
            if kind == "item":
                since.append(r)
            formula = _formula(key, label, refs, since)
            for col in ("C", "D"):
                cell = ws[f"{col}{r}"]
                cell.number_format = NUM
                if formula and cell.value is not None:
                    cell.value = "=" + formula.format(c=col)
                if kind in ("sub", "tot"):
                    cell.font = bold
                    cell.border = Border(top=thin, bottom=double if kind == "tot" else None)
            if kind in ("sub", "tot"):
                ws.cell(r, 1).font = bold
        ws.append([])
        ws.append([opts.get("sign_title", "")])
        ws.append([opts.get("sign_name", "")])

    # notes
    ws = wb.create_sheet("الإيضاحات")
    header(ws, "إيضاحات حول القوائم المالية")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 14, 46
    tables = DR.note_tables(dr, ctx, notes)
    last_kind = None
    for no, title, kind, rows, t1, t2 in tables:
        if kind != last_kind:
            ws.append([])
            ws.append(["إيضاحات قائمة الدخل" if kind == "is" else "إيضاحات قائمة المركز المالي"])
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=12, underline="single")
            last_kind = kind
        ws.append([])
        ws.append([f"إيضاح ({no})", title])
        ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = bold
        ws.append(["رقم الحساب", "اسم الحساب", labels["comp"] if t2 is not None else "", labels["cur"]])
        for c in ws[ws.max_row]:
            c.font = bold
            c.border = Border(bottom=thin)
        first = ws.max_row + 1
        for code, name, v1, v2 in rows:
            ws.append([code, name, v2, v1])
            ws.cell(ws.max_row, 3).number_format = ws.cell(ws.max_row, 4).number_format = NUM
        last = ws.max_row
        ws.append(["", "الإجمالي", f"=SUM(C{first}:C{last})" if t2 is not None else None, f"=SUM(D{first}:D{last})"])
        for col in (3, 4):
            c = ws.cell(ws.max_row, col)
            c.number_format, c.font, c.border = NUM, bold, Border(top=thin, bottom=double)
        ws.cell(ws.max_row, 2).font = bold
    if dr.checks:
        ws = wb.create_sheet("ملاحظات المراجعة")
        ws.sheet_view.rightToLeft = True
        ws.column_dimensions["A"].width = 120
        for c in dr.checks:
            ws.append([c])
    wb.save(path)
    return path


def _formula(key, label, refs, since):
    """Excel formula (with {c} for the column) for a total row, so the draft stays consistent when edited."""
    def ref(lbl):
        return f"{{c}}{refs[lbl]}" if lbl in refs else "0"
    if key == "fp":
        special = {"إجمالي الموجودات": ("إجمالي الموجودات المتداولة", "إجمالي الموجودات غير المتداولة"),
                   "إجمالي المطلوبات": ("إجمالي المطلوبات المتداولة", "إجمالي المطلوبات غير المتداولة"),
                   "إجمالي المطلوبات وحقوق الشركاء": ("إجمالي المطلوبات", "إجمالي حقوق الشركاء")}
        if label in special:
            return "+".join(ref(x) for x in special[label])
        if label.startswith("إجمالي") and since:
            return f"SUM({{c}}{since[0]}:{{c}}{since[-1]})"
        return None
    if key == "is":
        L = DR.IS_LABEL
        rules = {
            "مجمل الربح": f"{ref(L['sales'])}-{ref(L['cogs'])}",
            "إجمالي المصروفات": f"{ref(L['selling'])}+{ref(L['admin'])}",
            "صافي الدخل من العمليات الرئيسية": f"{ref('مجمل الربح')}-{ref('إجمالي المصروفات')}",
            "صافي الدخل بعد المصاريف التمويلية": f"{ref('صافي الدخل من العمليات الرئيسية')}-{ref(L['finance'])}",
            "صافي الربح قبل الزكاة": f"{ref('صافي الدخل بعد المصاريف التمويلية')}+{ref(L['invest'])}+{ref(L['other'])}",
            "صافي الدخل": f"{ref('صافي الربح قبل الزكاة')}-{ref(L['zakat'])}",
        }
        return rules.get(label)
    return None


# ------------------------------------------------------------------ PDF (HTML for QTextDocument)
CSS = """
body { font-size: 10.5pt; }
td { padding: 2px 4px; }
.num { text-align: center; }
"""


def _head(dr, opts, title, currency, logo_html):
    # Qt lays table columns left to right: the logo cell first (left), the company cell last (right)
    return (f'<table width="100%" cellspacing="0" cellpadding="0"><tr>'
            f'<td width="40%" align="left">{logo_html}</td>'
            f'<td width="60%" align="right" dir="rtl"><span style="font-size:10pt;">{html.escape(dr.company)}<br/>'
            f'{html.escape(opts.get("legal_form", ""))}</span></td></tr></table>'
            '<hr style="height:2px; background:#000;"/>'
            f'<p dir="rtl" align="right" style="font-size:11.5pt; margin:8px 0 0 0;">{html.escape(title)}</p>'
            f'<p dir="rtl" align="right" style="margin:0 0 10px 0;">{html.escape(currency)}</p>')


def _signature(opts):
    if not (opts.get("sign_title") or opts.get("sign_name")):
        return ""
    return ('<br/><br/><table width="100%"><tr><td width="40%" align="left" dir="rtl" style="font-size:11pt;">'
            f'<b>{html.escape(opts.get("sign_title", ""))}<br/>{html.escape(opts.get("sign_name", ""))}</b></td>'
            '<td width="60%"></td></tr></table>')


def _num_cell(v, kind, width):
    style = "font-weight:bold;" if kind in ("sub", "tot") else ""
    deco = "text-decoration: underline;" if kind == "tot" else ""
    return f'<td width="{width}" align="center" style="{style}{deco}">{money(v)}</td>'


def to_html(dr, ctx, opts, logo_path=""):
    labels = period_labels(dr, ctx)
    blocks, notes = DR.statements(dr, ctx)
    currency = opts.get("currency", "ريال سعودي")
    logo = f'<img src="{html.escape(logo_path)}" height="70"/>' if logo_path else ""
    parts = []
    for i, (key, rows) in enumerate(blocks):
        brk = ' style="page-break-before: always;"' if i else ""
        lab = period_labels(dr, ctx, key)
        body = [f'<div{brk}>', _head(dr, opts, TITLES[key].format(**lab), currency, logo),
                '<table width="100%" cellspacing="0" cellpadding="3">',
                # visual order left → right: current, comparative, note, label
                f'<tr><td width="15%" align="center"><b><u>{lab["cur"]}</u></b></td>'
                f'<td width="15%" align="center"><b><u>{lab["comp"]}</u></b></td>'
                '<td width="7%" align="center"><b>إيضاح</b></td><td width="63%"></td></tr>']
        for label, note, cur, comp, kind in rows:
            if kind == "blank":
                body.append('<tr><td colspan="4">&nbsp;</td></tr>')
                continue
            if kind == "hdr":
                body.append(f'<tr><td></td><td></td><td></td><td align="right" dir="rtl"><b><u>{html.escape(label)}</u></b></td></tr>')
                continue
            weight = "font-weight:bold;" if kind in ("sub", "tot") else ""
            body.append("<tr>" + _num_cell(cur, kind, "15%") + _num_cell(comp, kind, "15%")
                        + f'<td width="7%" align="center">{note}</td>'
                        f'<td width="63%" align="right" dir="rtl" style="{weight}">{html.escape(label)}</td></tr>')
        body.append("</table>")
        body.append(_signature(opts))
        body.append("</div>")
        parts.append("".join(body))

    # notes
    tables = DR.note_tables(dr, ctx, notes)
    out = ['<div style="page-break-before: always;">',
           _head(dr, opts, "إيضاحات حول القوائم المالية", currency, logo)]
    last_kind = None
    for no, title, kind, rows, t1, t2 in tables:
        if kind != last_kind:
            out.append(f'<p dir="rtl" align="right" style="font-size:11.5pt; margin-top:12px;"><b><u>'
                       f'{"إيضاحات قائمة الدخل" if kind == "is" else "إيضاحات قائمة المركز المالي"}</u></b></p>')
            last_kind = kind
        out.append(f'<p dir="rtl" align="right" style="margin:10px 0 2px 0;"><b>({no}) {html.escape(title)}</b></p>')
        out.append('<table width="100%" cellspacing="0" cellpadding="2" border="0">'
                   f'<tr><td width="17%" align="center"><b><u>{labels["cur"]}</u></b></td>'
                   f'<td width="17%" align="center"><b><u>{labels["comp"] if t2 is not None else ""}</u></b></td>'
                   '<td width="50%" align="right" dir="rtl"><b>اسم الحساب</b></td>'
                   '<td width="16%" align="right"><b>رقم الحساب</b></td></tr>')
        for code, name, v1, v2 in rows:
            out.append(f'<tr><td align="center">{money(v1)}</td><td align="center">{money(v2)}</td>'
                       f'<td align="right" dir="rtl">{html.escape(name)}</td><td align="right">{html.escape(code)}</td></tr>')
        out.append(f'<tr>{_num_cell(t1, "tot", "17%")}{_num_cell(t2, "tot", "17%")}'
                   '<td align="right" dir="rtl"><b>الإجمالي</b></td><td></td></tr></table>')
    out.append(_signature(opts))
    out.append("</div>")
    parts.append("".join(out))
    return f'<html><head><style>{CSS}</style></head><body dir="rtl">{"".join(parts)}</body></html>'


def to_pdf(dr, ctx, path, opts, logo_path=""):
    import sys

    from PySide6.QtCore import QMarginsF, QSizeF, Qt
    from PySide6.QtGui import QFont, QPageLayout, QPageSize, QPdfWriter, QTextDocument, QTextOption
    writer = QPdfWriter(path)
    writer.setPageLayout(QPageLayout(QPageSize(QPageSize.A4), QPageLayout.Portrait,
                                     QMarginsF(18, 14, 18, 14), QPageLayout.Millimeter))
    writer.setResolution(96)
    writer.setTitle(f"{dr.company} — مسودة القوائم المالية")
    doc = QTextDocument()
    doc.setDefaultFont(QFont("Segoe UI" if sys.platform == "win32" else "DejaVu Sans", 9))
    opt = QTextOption()
    opt.setTextDirection(Qt.RightToLeft)
    doc.setDefaultTextOption(opt)
    doc.setHtml(to_html(dr, ctx, opts, logo_path))
    rect = writer.pageLayout().paintRectPixels(writer.resolution())
    doc.setPageSize(QSizeF(rect.width(), rect.height()))
    doc.print_(writer)
    return path
