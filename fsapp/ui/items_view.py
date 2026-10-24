"""Items and customers: the customer master (code → type, rep, branch, region) and the items list with its
classification (main group → sub group 1 → sub group 2 → main item → item), guessed classes to review."""
import os

from PySide6.QtCore import QSettings, Qt, QUrl
from PySide6.QtGui import QColor, QDesktopServices
from PySide6.QtWidgets import (QAbstractItemView, QComboBox, QCompleter, QDialog, QFileDialog, QHBoxLayout,
                               QHeaderView, QLabel, QLineEdit, QMessageBox, QPushButton, QStyledItemDelegate,
                               QTableWidget, QTableWidgetItem, QTabWidget, QVBoxLayout, QWidget)

from .. import importer
from ..system import items as I
from .progress import Progress
from .table_tools import freeze

NAVY = "#12355b"
GUESS_BG = QColor("#fff4d6")
NONE_BG = QColor("#fde8e8")


def _info(parent, text):
    QMessageBox.information(parent, "تم", text)


def _error(parent, text):
    QMessageBox.warning(parent, "تنبيه", text)


class SuggestDelegate(QStyledItemDelegate):
    """Free text with suggestions from the values already in the column."""

    def __init__(self, table, column):
        super().__init__(table)
        self.table, self.column = table, column

    def createEditor(self, parent, option, index):
        e = QLineEdit(parent)
        values = sorted({self.table.item(r, self.column).text() for r in range(self.table.rowCount())
                         if self.table.item(r, self.column) and self.table.item(r, self.column).text()})
        c = QCompleter(values, e)
        c.setCaseSensitivity(Qt.CaseInsensitive)
        c.setFilterMode(Qt.MatchContains)
        e.setCompleter(c)
        return e


def _bar(*buttons):
    h = QHBoxLayout()
    for text, fn in buttons:
        if text is None:
            h.addStretch()
            continue
        b = QPushButton(text)
        b.setMinimumHeight(32)
        b.clicked.connect(fn)
        h.addWidget(b)
    return h


def _note(text):
    h = QLabel(text)
    h.setWordWrap(True)
    h.setStyleSheet("padding:6px; background:#eef4fb; border-radius:6px;")
    return h


def _save_xlsx(parent, name, headers, rows):
    start = QSettings().value("last_dir", os.path.expanduser("~"))
    path, _ = QFileDialog.getSaveFileName(parent, "تصدير", os.path.join(start, name), "Excel (*.xlsx)")
    if not path:
        return
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="12355B")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(45, len(h) + 6))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(path)
    except OSError as e:
        _error(parent, f"تعذر الحفظ (الملف مفتوح؟):\n{e}")
        return
    QDesktopServices.openUrl(QUrl.fromLocalFile(path))


# ================================================================== customers
class CustomersTab(QWidget):
    COLS = [("code", "رقم العميل"), ("name", "اسم العميل"), ("ctype", "النوع"), ("rep", "المندوب"),
            ("branch", "الفرع"), ("region", "المنطقة")]

    def __init__(self, win):
        super().__init__()
        self.win = win
        lay = QVBoxLayout(self)
        lay.addWidget(_note("بيانات العملاء تربط مبيعات الأصناف بالمندوب والفرع والمنطقة برقم العميل. ارفعها من ملف Excel "
                            "(رقم العميل، الاسم، النوع، المندوب، الفرع، المنطقة) أو عدّل هنا. العميل غير الموجود هنا "
                            "تؤخذ بياناته من آخر أعمار عملاء مرفوعة، والمنطقة الفارغة تؤخذ من الفرع في الإعدادات."))
        lay.addLayout(_bar(("📥 رفع ملف بيانات العملاء", self.import_file), ("＋ إضافة عميل", self.add),
                           ("− حذف المحدد", self.remove), ("📤 تصدير Excel", self.export), (None, None),
                           ("💾 حفظ", self.save)))
        top = QHBoxLayout()
        top.addWidget(QLabel("بحث:"))
        self.search = QLineEdit()
        self.search.setPlaceholderText("رقم أو اسم أو مندوب ...")
        self.search.textChanged.connect(self.filter)
        top.addWidget(self.search, 1)
        self.status = QLabel("")
        top.addWidget(self.status)
        lay.addLayout(top)
        self.table = QTableWidget(0, len(self.COLS))
        self.table.setHorizontalHeaderLabels([h for _, h in self.COLS])
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.verticalHeader().setVisible(False)
        for c in range(2, 6):
            self.table.setItemDelegateForColumn(c, SuggestDelegate(self.table, c))
        lay.addWidget(self.table, 1)
        freeze(self.table, 2)

    def load(self):
        rows = list(self.win.db.customer_master().values())
        t = self.table
        t.setSortingEnabled(False)
        t.setRowCount(len(rows))
        for r, rec in enumerate(rows):
            for c, (k, _) in enumerate(self.COLS):
                t.setItem(r, c, QTableWidgetItem(rec.get(k) or ""))
        t.setSortingEnabled(True)
        t.resizeColumnsToContents()
        self.status.setText(f"{len(rows)} عميل")
        self.filter()

    def filter(self):
        q = I.norm(self.search.text())
        for r in range(self.table.rowCount()):
            text = " ".join(I.norm(self.table.item(r, c).text()) for c in range(self.table.columnCount())
                            if self.table.item(r, c))
            self.table.setRowHidden(r, bool(q) and q not in text)
        self.table._frozen.refresh()

    def rows(self):
        out = []
        for r in range(self.table.rowCount()):
            rec = {k: (self.table.item(r, c).text().strip() if self.table.item(r, c) else "")
                   for c, (k, _) in enumerate(self.COLS)}
            if rec["code"]:
                out.append(rec)
        return out

    def add(self):
        self.search.clear()
        t = self.table
        t.setSortingEnabled(False)
        t.insertRow(0)
        for c in range(t.columnCount()):
            t.setItem(0, c, QTableWidgetItem(""))
        t.setCurrentCell(0, 0)
        t.editItem(t.item(0, 0))

    def remove(self):
        for r in sorted({i.row() for i in self.table.selectedIndexes()}, reverse=True):
            self.table.removeRow(r)

    def save(self):
        rows = self.rows()
        codes = [r["code"] for r in rows]
        dup = {c for c in codes if codes.count(c) > 1}
        if dup:
            _error(self, "أرقام عملاء مكررة: " + "، ".join(sorted(dup)[:10]))
            return
        self.win.db.save_customer_master(rows)
        self.win.data_changed(reload_pages=False)
        self.status.setText(f"{len(rows)} عميل — تم الحفظ")

    def export(self):
        _save_xlsx(self, "بيانات العملاء.xlsx", [h for _, h in self.COLS],
                   [[r[k] for k, _ in self.COLS] for r in self.rows()])

    def import_file(self):
        from .system_window import ColumnMapDialog, ROLE_LABELS
        start = QSettings().value("last_dir", os.path.expanduser("~"))
        path, _ = QFileDialog.getOpenFileName(self, "ملف بيانات العملاء", start, "Excel / CSV (*.xlsx *.xlsm *.csv *.txt)")
        if not path:
            return
        QSettings().setValue("last_dir", os.path.dirname(path))
        try:
            with Progress(self, f"قراءة {os.path.basename(path)}") as bar:
                bar.step(10, "جاري القراءة")
                rows = importer.read_table(path)
        except Exception as e:  # noqa: BLE001
            _error(self, f"تعذر قراءة الملف:\n{e}")
            return
        first, mapping = I.guess_master(rows)
        roles = ["m_" + k for k, _ in self.COLS]
        for k, h in self.COLS:
            ROLE_LABELS.setdefault("m_" + k, h)
        cm = ColumnMapDialog(self, rows, roles, first, {"m_" + k: v for k, v in mapping.items()},
                             "حدد الأعمدة. يكفي رقم العميل مع المندوب أو الفرع.")
        if cm.exec() != QDialog.Accepted:
            return
        m = {k[2:]: v for k, v in cm.mapping().items()}
        if "code" not in m:
            _error(self, "لازم تحدد عمود رقم العميل.")
            return
        recs = I.read_master(rows, cm.first_row(), m)
        if not recs:
            _error(self, "لم أجد عملاء في الملف.")
            return
        existing = self.win.db.customer_master()
        replace = True
        if existing:
            box = QMessageBox(self)
            box.setWindowTitle("بيانات العملاء")
            box.setText(f"في الملف {len(recs)} عميل، والمحفوظ حاليًا {len(existing)} عميل.")
            b_rep = box.addButton("استبدال الكل بالملف", QMessageBox.AcceptRole)
            b_upd = box.addButton("تحديث وإضافة فقط (إبقاء غير الموجودين في الملف)", QMessageBox.AcceptRole)
            box.addButton("إلغاء", QMessageBox.RejectRole)
            box.exec()
            if box.clickedButton() not in (b_rep, b_upd):
                return
            replace = box.clickedButton() is b_rep
            if not replace:   # keep the saved fields the file doesn't have
                for r in recs:
                    old = existing.get(r["code"], {})
                    for k, _ in self.COLS:
                        if not r.get(k) and old.get(k):
                            r[k] = old[k]
        self.win.db.save_customer_master(recs, replace=replace)
        self.win.data_changed(reload_pages=False)
        self.load()
        _info(self, f"تم حفظ بيانات {len(recs)} عميل.")


# ================================================================== items
class ItemsTab(QWidget):
    COLS = [("code", "رقم الصنف"), ("name", "اسم الصنف"), ("unit", "الوحدة"), ("grp_main", "المجموعة الرئيسية"),
            ("grp_sub1", "المجموعة الفرعية 1"), ("grp_sub2", "المجموعة الفرعية 2"), ("main_item", "الصنف الرئيسي"),
            ("itype", "نوع الصنف"), ("status", "الحالة"), ("guess_from", "التخمين من")]
    EDIT = {"grp_main", "grp_sub1", "grp_sub2", "main_item", "itype"}

    def __init__(self, win):
        super().__init__()
        self.win = win
        self.data = {}
        self.dirty = set()
        lay = QVBoxLayout(self)
        lay.addWidget(_note("التصنيف: المجموعة الرئيسية ← الفرعية 1 ← الفرعية 2 ← الصنف الرئيسي ← الصنف. ارفع ملف "
                            "الأصناف (ورقة «الأصناف» من نموذج طلب المشتريات). الأصناف غير المصنفة يُخمَّن تصنيفها من أقرب "
                            "صنف مصنف بالاسم (بالأصفر للمراجعة)، وما لا يشبه شيئًا يبقى «غير مصنف» (بالأحمر). عدّل أي "
                            "خانة مباشرة، أو صدّر للمراجعة في Excel وارفع الملف بعد التعديل."))
        lay.addLayout(_bar(("📥 رفع ملف الأصناف", self.import_file), ("🔮 تخمين غير المصنف", self.reguess),
                           ("✔ اعتماد تخمين المحدد", self.accept_selected), ("📤 تصدير للمراجعة", self.export),
                           (None, None), ("💾 حفظ التعديلات", self.save)))
        top = QHBoxLayout()
        top.addWidget(QLabel("عرض:"))
        self.view = QComboBox()
        for text, key in (("الكل", "all"), ("التخمين (للمراجعة)", "guessed"), ("تخمين ضعيف (أقل من 60%)", "weak"),
                          ("غير مصنف", "none"), ("مصنف", "ok")):
            self.view.addItem(text, key)
        self.view.currentIndexChanged.connect(self.filter)
        top.addWidget(self.view)
        top.addWidget(QLabel("بحث:"))
        self.search = QLineEdit()
        self.search.setPlaceholderText("رقم أو اسم أو مجموعة ...")
        self.search.textChanged.connect(self.filter)
        top.addWidget(self.search, 1)
        self.status = QLabel("")
        top.addWidget(self.status)
        lay.addLayout(top)
        self.table = QTableWidget(0, len(self.COLS))
        self.table.setHorizontalHeaderLabels([h for _, h in self.COLS])
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.verticalHeader().setVisible(False)
        for c, (k, _) in enumerate(self.COLS):
            if k in self.EDIT:
                self.table.setItemDelegateForColumn(c, SuggestDelegate(self.table, c))
        self.table.itemChanged.connect(self._edited)
        lay.addWidget(self.table, 1)
        freeze(self.table, 2)

    def _status_text(self, rec):
        if not I.is_classified(rec):
            return "غير مصنف"
        if rec.get("guessed"):
            return f"تخمين {float(rec.get('guess_score') or 0):.0%}"
        return "مصنف"

    def load(self):
        self.data = self.win.db.items()
        self.dirty = set()
        t = self.table
        t.blockSignals(True)
        t.setSortingEnabled(False)
        t.setRowCount(len(self.data))
        for r, rec in enumerate(self.data.values()):
            self._fill_row(r, rec)
        t.setSortingEnabled(True)
        t.blockSignals(False)
        t.resizeColumnsToContents()
        t.setColumnWidth(1, min(t.columnWidth(1), 380))
        self.filter()

    def _fill_row(self, r, rec):
        t = self.table
        st = self._status_text(rec)
        bg = NONE_BG if st == "غير مصنف" else GUESS_BG if rec.get("guessed") else None
        for c, (k, _) in enumerate(self.COLS):
            v = st if k == "status" else (rec.get(k) or "")
            if k in I.LEVELS and not v:
                v = ""
            it = QTableWidgetItem(str(v))
            if k not in self.EDIT:
                it.setFlags(it.flags() & ~Qt.ItemIsEditable)
            if k == "code":
                it.setData(Qt.UserRole, rec["code"])
            if bg is not None:
                it.setBackground(bg)
            t.setItem(r, c, it)

    def _edited(self, it):
        k = self.COLS[it.column()][0]
        if k not in self.EDIT:
            return
        code = self.table.item(it.row(), 0).data(Qt.UserRole)
        rec = self.data.get(code)
        if rec is None:
            return
        rec[k] = it.text().strip()
        if k in I.LEVELS:
            rec["guessed"], rec["guess_from"], rec["guess_score"] = 0, "", 0.0   # reviewed by hand
        self.dirty.add(code)
        self.table.blockSignals(True)
        sort = self.table.isSortingEnabled()
        self.table.setSortingEnabled(False)
        self._fill_row(it.row(), rec)
        self.table.setSortingEnabled(sort)
        self.table.blockSignals(False)
        self._count()

    def filter(self):
        q = I.norm(self.search.text())
        mode = self.view.currentData()
        for r in range(self.table.rowCount()):
            rec = self.data.get(self.table.item(r, 0).data(Qt.UserRole), {})
            ok = {"all": True, "guessed": bool(rec.get("guessed")),
                  "weak": bool(rec.get("guessed")) and float(rec.get("guess_score") or 0) < 0.6,
                  "none": not I.is_classified(rec),
                  "ok": I.is_classified(rec) and not rec.get("guessed")}[mode]
            if ok and q:
                text = " ".join(I.norm(rec.get(k, "")) for k in ("code", "name", *I.LEVELS))
                ok = q in text
            self.table.setRowHidden(r, not ok)
        self.table._frozen.refresh()
        self._count()

    def _count(self):
        vals = list(self.data.values())
        n_ok = sum(1 for r in vals if I.is_classified(r) and not r.get("guessed"))
        n_g = sum(1 for r in vals if r.get("guessed"))
        n_none = sum(1 for r in vals if not I.is_classified(r))
        self.status.setText(f"{len(vals)} صنف: مصنف {n_ok} — تخمين {n_g} — غير مصنف {n_none}"
                            + (f" — ⚠ {len(self.dirty)} تعديل غير محفوظ" if self.dirty else ""))

    def selected_codes(self):
        return [self.table.item(r, 0).data(Qt.UserRole) for r in sorted({i.row() for i in self.table.selectedIndexes()})]

    def accept_selected(self):
        codes = [c for c in self.selected_codes() if self.data.get(c, {}).get("guessed")]
        if not codes:
            _info(self, "حدد أسطر أصناف مخمَّنة (بالأصفر) ثم اضغط اعتماد.")
            return
        for c in codes:
            self.data[c]["guessed"] = 0
            self.dirty.add(c)
        self.save(silent=True)
        _info(self, f"تم اعتماد تصنيف {len(codes)} صنف.")

    def save(self, silent=False):
        if not self.dirty:
            if not silent:
                _info(self, "لا توجد تعديلات.")
            return
        db = self.win.db
        for code in self.dirty:
            rec = self.data[code]
            db.update_item(code, **{k: rec.get(k, "") for k in (*I.LEVELS, "itype", "guessed", "guess_from",
                                                                 "guess_score")})
        n = len(self.dirty)
        self.win.data_changed(reload_pages=False)
        self.load()
        if not silent:
            _info(self, f"تم حفظ تعديل {n} صنف.")

    def reguess(self):
        vals = list(self.data.values())
        if not vals:
            _info(self, "ارفع ملف الأصناف أولًا.")
            return
        if self.dirty:
            self.save(silent=True)
            vals = list(self.data.values())
        with Progress(self, "تخمين تصنيف الأصناف") as bar:
            n = I.guess_classes(vals, progress=bar)
            bar.step(95, "حفظ")
            self.win.db.save_items(vals)
            self.win.data_changed(reload_pages=False)
            self.load()
        _info(self, f"تم تخمين تصنيف {n} صنف — راجعها من «عرض: التخمين».")

    def export(self):
        heads = [h for _, h in self.COLS] + ["درجة التشابه"]
        rows = []
        for rec in self.data.values():
            rows.append([rec.get(k) or (I.UNCLASSIFIED if k in I.LEVELS else "") if k != "status"
                         else self._status_text(rec) for k, _ in self.COLS] + [rec.get("guess_score") or None])
        _save_xlsx(self, "الأصناف وتصنيفها.xlsx", heads, rows)

    def import_file(self):
        start = QSettings().value("last_dir", os.path.expanduser("~"))
        path, _ = QFileDialog.getOpenFileName(self, "ملف الأصناف (ورقة «الأصناف»)", start, "Excel (*.xlsx *.xlsm)")
        if not path:
            return
        QSettings().setValue("last_dir", os.path.dirname(path))
        bar = Progress(self, "رفع ملف الأصناف")
        try:
            import openpyxl
            bar.step(5, f"فتح {os.path.basename(path)}")
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            names = wb.sheetnames
            pick = [n for n in names if I.norm(n) == I.norm("الأصناف")] or names
            bar.step(20, "قراءة ورقة الأصناف")
            sheets = [(n, [list(r) for r in wb[n].iter_rows(values_only=True)]) for n in pick[:6]]
            wb.close()
            recs = I.read_items_sheets(sheets)
        except Exception as e:  # noqa: BLE001
            bar.close()
            _error(self, f"تعذر قراءة ملف الأصناف:\n{e}")
            return
        # keep classifications set here by hand for items that are still unclassified in the file
        old = self.win.db.items()
        kept = 0
        for r in recs:
            o = old.get(r["code"])
            if not I.is_classified(r) and o and I.is_classified(o) and not o.get("guessed"):
                for lv in I.LEVELS:
                    r[lv] = o.get(lv, "")
                kept += 1
        before = sum(1 for r in recs if I.is_classified(r))
        bar.title = "رفع ملف الأصناف — تخمين التصنيف"
        n = I.guess_classes(recs, progress=bar)
        bar.step(95, f"حفظ {len(recs)} صنف")
        self.win.db.save_items(recs)
        self.win.data_changed(reload_pages=False)
        self.load()
        bar.close()
        none = sum(1 for r in recs if not I.is_classified(r))
        msg = (f"تم حفظ {len(recs)} صنف: مصنف في الملف {before - kept}"
               + (f"، ومحفوظ تصنيفه من مراجعة سابقة {kept}" if kept else "")
               + f"، وتخمين {n} (بالأصفر للمراجعة)، وغير مصنف {none}.")
        self.view.setCurrentIndex(1 if n else 0)
        _info(self, msg)


class ItemsPage(QWidget):
    def __init__(self, win):
        super().__init__()
        self.win = win
        lay = QVBoxLayout(self)
        self.tabs = QTabWidget()
        self.customers = CustomersTab(win)
        self.items = ItemsTab(win)
        self.tabs.addTab(self.items, "الأصناف وتصنيفها")
        self.tabs.addTab(self.customers, "بيانات العملاء")
        lay.addWidget(self.tabs)

    def load(self):
        self.items.load()
        self.customers.load()
