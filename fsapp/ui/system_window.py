"""Main window of the multi-company financial reporting system."""
import os
import sys
import traceback

from PySide6.QtCore import QMarginsF, QSettings, QSizeF, Qt
from PySide6.QtCore import QUrl
from PySide6.QtGui import (QAction, QColor, QDesktopServices, QFont, QPageLayout, QPageSize, QPdfWriter, QTextDocument,
                           QTextOption)
from PySide6.QtWidgets import (
    QAbstractItemView, QApplication, QButtonGroup, QCheckBox, QComboBox, QDialog, QDialogButtonBox,
    QDoubleSpinBox, QFileDialog, QFormLayout, QGridLayout, QGroupBox, QHBoxLayout, QHeaderView, QLabel,
    QLineEdit, QListWidget, QListWidgetItem, QMainWindow, QMessageBox, QPlainTextEdit, QPushButton,
    QRadioButton, QScrollArea, QSpinBox, QSplitter, QStackedWidget, QStyledItemDelegate, QTableWidget, QTableWidgetItem,
    QTabWidget, QTextBrowser, QToolBar, QVBoxLayout, QWidget,
)

from .. import importer
from ..system import aging, calc, defaults, diagnose, migrate, render
from ..system import reports as R
from ..system.db import ACCOUNT_FIELDS, Account, CompanyDB
from ..updater import UpdateChecker
from ..version import BUILD
from .aging_view import AgingPage
from .dashboard_view import DashboardPage
from . import zoom
from .table_tools import freeze
from .zoom import pt

NAVY = "#12355b"


def col_letter(i):
    s, i = "", i + 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def info(parent, text):
    QMessageBox.information(parent, "معلومة", text)


def error(parent, text):
    QMessageBox.critical(parent, "خطأ", text)


def write_pdf(html, path, landscape=False, rtl=True):
    writer = QPdfWriter(path)
    writer.setPageLayout(QPageLayout(QPageSize(QPageSize.A4),
                                     QPageLayout.Landscape if landscape else QPageLayout.Portrait,
                                     QMarginsF(10, 10, 10, 10), QPageLayout.Millimeter))
    writer.setResolution(96)
    doc = QTextDocument()
    doc.setDefaultFont(QFont("Segoe UI" if sys.platform == "win32" else "DejaVu Sans", 8))
    opt = QTextOption()
    opt.setTextDirection(Qt.RightToLeft if rtl else Qt.LeftToRight)
    doc.setDefaultTextOption(opt)
    doc.setHtml(html)
    rect = writer.pageLayout().paintRectPixels(writer.resolution())
    doc.setPageSize(QSizeF(rect.width(), rect.height()))
    doc.print_(writer)


# ================================================================== small dialogs
class MonthDialog(QDialog):
    def __init__(self, parent, title, default):
        super().__init__(parent)
        self.setWindowTitle(title)
        f = QFormLayout(self)
        self.year = QSpinBox()
        self.year.setRange(2000, 2100)
        self.year.setValue(int(default[:4]))
        self.month = QComboBox()
        self.month.addItems(calc.MONTHS_AR)
        self.month.setCurrentIndex(int(default[5:7]) - 1)
        f.addRow("السنة:", self.year)
        f.addRow("الشهر:", self.month)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        f.addRow(bb)

    def period(self):
        return f"{self.year.value():04d}-{self.month.currentIndex() + 1:02d}"


ROLE_LABELS = {
    "code": "رقم الحساب", "name": "اسم الحساب", "branch": "الفرع",
    "od": "افتتاحي مدين", "oc": "افتتاحي دائن", "md": "حركة مدين", "mc": "حركة دائن",
    "cd": "ختامي مدين", "cc": "ختامي دائن",
    "no": "الرقم", "aname": "الاسم", "region": "المنطقة", "rep": "المندوب", "cls": "التصنيف",
    "ctype": "النوع (آجل / عهدة / نقدي)", "b0": "0 - 30", "b1": "31 - 60", "b2": "61 - 90", "b3": "91 - 120",
    "b4": "أكثر من 120", "guarantee": "ضمان بنكي", "account": "رقم الحساب", "credit": "فترة الائتمان (يوم)",
    "currency": "العملة",
}

AGING_ROLES = {
    "customers": ["no", "aname", "region", "branch", "rep", "cls", "ctype", "b0", "b1", "b2", "b3", "b4", "guarantee"],
    "suppliers": ["no", "aname", "account", "credit", "currency", "b0", "b1", "b2", "b3", "b4"],
}
AGING_WORDS = {
    "no": ("رقم العميل", "رقم المورد", "كود العميل", "كود المورد", "الرقم", "رقم"),
    "aname": ("اسم العميل", "اسم المورد", "الاسم", "اسم"),
    "region": ("المنطقة",), "branch": ("الفرع",), "rep": ("المندوب", "البائع"), "cls": ("التصنيف", "تصنيف"),
    "ctype": ("نوع العميل", "النوع", "نوع"), "b0": ("0 - 30", "0-30", "حالي"), "b1": ("31 - 60", "31-60"),
    "b2": ("61 - 90", "61-90"), "b3": ("91 - 120", "91-120"), "b4": ("أكثر من 120", "اكثر من 120", "120+", "> 120"),
    "guarantee": ("ضمان",), "account": ("رقم الحساب", "الحساب"), "credit": ("فترة الائتمان", "الائتمان", "مدة"),
    "currency": ("العملة",),
}


def guess_aging(rows, roles):
    for i, r in enumerate(rows[:30]):
        cells = [str(v).strip() if v is not None else "" for v in r]
        found = {}
        for role in roles:
            for w in AGING_WORDS[role]:
                c = next((k for k, text in enumerate(cells) if text and w in text and k not in found.values()), None)
                if c is not None:
                    found[role] = c
                    break
        if len(found) >= 4:
            return i + 1, found
    return 0, {}


class ColumnMapDialog(QDialog):
    """Pick which column holds each field of an imported table."""

    def __init__(self, parent, rows, roles, first, mapping, help_text):
        super().__init__(parent)
        self.rows = rows
        self.setWindowTitle("تحديد الأعمدة")
        self.resize(1050, 640)
        lay = QVBoxLayout(self)
        h = QLabel(help_text)
        h.setWordWrap(True)
        lay.addWidget(h)
        grid = QGridLayout()
        width = len(rows[0]) if rows else 0
        self.combos = {}
        for i, role in enumerate(roles):
            cb = QComboBox()
            cb.addItem("— غير موجود —", None)
            for c in range(width):
                hint = self._hint(c, first)
                cb.addItem(f"عمود {col_letter(c)}" + (f" — {hint}" if hint else ""), c)
            if mapping.get(role) is not None:
                cb.setCurrentIndex(mapping[role] + 1)
            self.combos[role] = cb
            grid.addWidget(QLabel(ROLE_LABELS[role]), i // 4 * 2, i % 4)
            grid.addWidget(cb, i // 4 * 2 + 1, i % 4)
        lay.addLayout(grid)
        row = QHBoxLayout()
        row.addWidget(QLabel("أول صف بيانات:"))
        self.first = QSpinBox()
        self.first.setRange(1, max(1, len(rows)))
        self.first.setValue(first + 1)
        row.addWidget(self.first)
        row.addStretch()
        lay.addLayout(row)
        t = QTableWidget(min(len(rows), 50), width)
        t.setHorizontalHeaderLabels([col_letter(c) for c in range(width)])
        t.setEditTriggers(QAbstractItemView.NoEditTriggers)
        for r in range(t.rowCount()):
            for c in range(width):
                v = rows[r][c]
                t.setItem(r, c, QTableWidgetItem("" if v is None else str(v)))
        t.resizeColumnsToContents()
        lay.addWidget(t, 1)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)

    def _hint(self, c, first):
        parts = [str(self.rows[r][c]).strip() for r in range(max(0, first - 3), first)
                 if self.rows[r][c] is not None and str(self.rows[r][c]).strip()]
        return " / ".join(parts)[:40]

    def mapping(self):
        return {r: cb.currentData() for r, cb in self.combos.items() if cb.currentData() is not None}

    def first_row(self):
        return self.first.value() - 1


class ListDelegate(QStyledItemDelegate):
    """Combo-box editor whose options come from a callable."""

    def __init__(self, parent, options):
        super().__init__(parent)
        self.options = options

    def createEditor(self, parent, option, index):
        cb = QComboBox(parent)
        cb.addItem("")
        cb.addItems(self.options())
        cb.setMaxVisibleItems(25)
        cb.activated.connect(lambda _=None, e=cb: (self.commitData.emit(e), self.closeEditor.emit(e)))
        return cb

    def setEditorData(self, editor, index):
        i = editor.findText(index.data() or "")
        editor.setCurrentIndex(max(0, i))
        editor.showPopup()

    def setModelData(self, editor, model, index):
        model.setData(index, editor.currentText())


class CodeDelegate(QStyledItemDelegate):
    """Combo editor showing a label per stored code (used for section / role / activity...)."""

    def __init__(self, parent, mapping):
        super().__init__(parent)
        self.mapping = mapping  # {code: label}

    def createEditor(self, parent, option, index):
        cb = QComboBox(parent)
        for code, label in self.mapping.items():
            cb.addItem(label, code)
        cb.activated.connect(lambda _=None, e=cb: (self.commitData.emit(e), self.closeEditor.emit(e)))
        return cb

    def setEditorData(self, editor, index):
        i = editor.findData(index.data(Qt.UserRole))
        editor.setCurrentIndex(max(0, i))
        editor.showPopup()

    def setModelData(self, editor, model, index):
        model.setData(index, editor.currentData(), Qt.UserRole)
        model.setData(index, editor.currentText())


def ro_item(text):
    it = QTableWidgetItem("" if text is None else str(text))
    it.setFlags(it.flags() & ~Qt.ItemIsEditable)
    return it


def code_item(code, mapping):
    it = QTableWidgetItem(mapping.get(code, code or ""))
    it.setData(Qt.UserRole, code)
    return it


# ================================================================== settings page
class SettingsPage(QWidget):
    def __init__(self, win):
        super().__init__()
        self.win = win
        lay = QVBoxLayout(self)
        box = QGroupBox("بيانات الشركة")
        f = QFormLayout(box)
        self.company = QLineEdit()
        self.currency = QLineEdit()
        self.fiscal = QComboBox()
        self.fiscal.addItems(calc.MONTHS_AR)
        self.credit = QDoubleSpinBox()
        self.credit.setRange(0, 100)
        self.credit.setSuffix(" %")
        self.alert_pct = QDoubleSpinBox()
        self.alert_pct.setRange(0, 1000)
        self.alert_pct.setSuffix(" %")
        self.alert_amt = QDoubleSpinBox()
        self.alert_amt.setRange(0, 1e12)
        self.alert_amt.setDecimals(0)
        self.alert_amt.setGroupSeparatorShown(True)
        self.admin = QLineEdit()
        f.addRow("اسم الشركة:", self.company)
        f.addRow("العملة:", self.currency)
        f.addRow("أول شهر في السنة المالية:", self.fiscal)
        f.addRow("نسبة المبيعات الآجلة من صافي المبيعات:", self.credit)
        f.addRow("حد التنبيه لتغير المصروف (نسبة):", self.alert_pct)
        f.addRow("حد التنبيه لتغير المصروف (مبلغ):", self.alert_amt)
        f.addRow("اسم مركز التكلفة غير الموزع (الإدارة):", self.admin)
        lay.addWidget(box)

        h = QHBoxLayout()
        bb = QGroupBox("الفروع (بنفس الاسم المستخدم في ملفات حركة الفروع)")
        bl = QVBoxLayout(bb)
        self.branches = QTableWidget(0, 2)
        self.branches.setHorizontalHeaderLabels(["الفرع", "المنطقة"])
        self.branches.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        bl.addWidget(self.branches)
        bl.addLayout(self._row_buttons(self.branches))
        cb = QGroupBox("العملات وسعر الصرف (آخر الفترة)")
        cl = QVBoxLayout(cb)
        self.curr = QTableWidget(0, 2)
        self.curr.setHorizontalHeaderLabels(["العملة", "سعر الصرف"])
        self.curr.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        cl.addWidget(self.curr)
        cl.addLayout(self._row_buttons(self.curr))
        h.addWidget(bb, 2)
        h.addWidget(cb, 1)
        pb = QGroupBox("قاعدة أرقام الحسابات (لتشخيص الخلل)")
        pl = QFormLayout(pb)
        self.prefix = {}
        for d in "123456789":
            combo = QComboBox()
            for code, name in diagnose.TYPE_NAMES.items():
                combo.addItem(name, code)
            self.prefix[d] = combo
            pl.addRow(f"حساب يبدأ بـ {d}:", combo)
        h.addWidget(pb, 1)
        lay.addLayout(h, 1)
        save = QPushButton("حفظ الإعدادات")
        save.setMinimumHeight(36)
        save.setStyleSheet(f"font-weight:700; background:{NAVY}; color:white;")
        save.clicked.connect(self.save)
        lay.addWidget(save)

    def _row_buttons(self, table):
        h = QHBoxLayout()
        add = QPushButton("＋ إضافة")
        add.clicked.connect(lambda: table.insertRow(table.rowCount()))
        rm = QPushButton("− حذف المحدد")
        rm.clicked.connect(lambda: [table.removeRow(r) for r in sorted({i.row() for i in table.selectedIndexes()}, reverse=True)])
        h.addWidget(add)
        h.addWidget(rm)
        h.addStretch()
        return h

    def load(self):
        db = self.win.db
        self.company.setText(db.setting("company") or "")
        self.currency.setText(db.setting("currency") or "")
        self.fiscal.setCurrentIndex(int(db.setting_float("fiscal_start")) - 1)
        self.credit.setValue(db.setting_float("credit_ratio") * 100)
        self.alert_pct.setValue(db.setting_float("alert_pct") * 100)
        self.alert_amt.setValue(db.setting_float("alert_amount"))
        self.admin.setText(db.setting("admin_label") or "")
        pm = diagnose.prefix_map(db)
        for d, combo in self.prefix.items():
            combo.setCurrentIndex(max(0, combo.findData(pm.get(d, ""))))
        for table, rows in ((self.branches, db.branch_rows()), (self.curr, db.currencies())):
            table.setRowCount(len(rows))
            for r, row in enumerate(rows):
                for c, v in enumerate(row):
                    table.setItem(r, c, QTableWidgetItem("" if v is None else str(v)))

    def save(self):
        db = self.win.db
        db.set_setting("company", self.company.text().strip())
        db.set_setting("currency", self.currency.text().strip())
        db.set_setting("fiscal_start", self.fiscal.currentIndex() + 1)
        db.set_setting("credit_ratio", self.credit.value() / 100)
        db.set_setting("alert_pct", self.alert_pct.value() / 100)
        db.set_setting("alert_amount", self.alert_amt.value())
        db.set_setting("admin_label", self.admin.text().strip() or defaults.SETTINGS["admin_label"])
        import json
        db.set_setting("prefix_map", json.dumps({d: c.currentData() for d, c in self.prefix.items()}))

        def cells(t):
            return [[(t.item(r, c).text() if t.item(r, c) else "") for c in range(t.columnCount())]
                    for r in range(t.rowCount())]

        db.save_branches([(a, b) for a, b in cells(self.branches)])
        rates = []
        for code, rate in cells(self.curr):
            try:
                rates.append((code, float(rate.replace(",", "")) if rate.strip() else None))
            except ValueError:
                rates.append((code, None))
        db.save_currencies(rates)
        self.win.data_changed()
        info(self, "تم حفظ الإعدادات.")


# ================================================================== lists page
class ListsPage(QWidget):
    TABS = [
        ("fs_items", "بنود القوائم المالية", ["البند", "القسم", "دور في النسب"],
         [defaults.SECTION_NAMES, defaults.ROLE_NAMES]),
        ("cf_items", "بنود التدفقات النقدية", ["البند", "النشاط", "مجموعة «الكاش راح فين»"],
         [defaults.ACTIVITY_NAMES, {"": "—", **defaults.BUCKET_NAMES}]),
        ("branch_cats", "فئات دخل الفروع", ["الفئة", "النوع"], [defaults.KIND_NAMES]),
        ("exp_groups", "مجموعات المصروفات", ["المجموعة"], []),
    ]

    def __init__(self, win):
        super().__init__()
        self.win = win
        lay = QVBoxLayout(self)
        h = QLabel("هذه القوائم تظهر كاختيارات في شاشة ربط الحسابات. ترتيب البنود هنا هو ترتيبها في التقارير. "
                   "تغيير اسم بند يُحدّث الحسابات المربوطة به تلقائيًا.")
        h.setWordWrap(True)
        h.setStyleSheet("padding:6px; background:#eef4fb; border-radius:6px;")
        lay.addWidget(h)
        self.tabs = QTabWidget()
        self.tables = {}
        for key, title, headers, maps in self.TABS:
            w = QWidget()
            v = QVBoxLayout(w)
            t = QTableWidget(0, len(headers))
            t.setHorizontalHeaderLabels(headers)
            t.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
            t.setSelectionBehavior(QAbstractItemView.SelectRows)
            for i, m in enumerate(maps, start=1):
                t.setItemDelegateForColumn(i, CodeDelegate(t, m))
            v.addWidget(t)
            bar = QHBoxLayout()
            for text, fn in (("＋ إضافة", self.add), ("− حذف", self.remove), ("▲ لأعلى", lambda k=key: self.move(-1)),
                             ("▼ لأسفل", lambda k=key: self.move(1))):
                b = QPushButton(text)
                b.clicked.connect(fn)
                bar.addWidget(b)
            bar.addStretch()
            save = QPushButton("حفظ")
            save.setStyleSheet(f"font-weight:700; background:{NAVY}; color:white; padding:4px 18px;")
            save.clicked.connect(self.save)
            bar.addWidget(save)
            v.addLayout(bar)
            self.tables[key] = (t, maps)
            self.tabs.addTab(w, title)
        lay.addWidget(self.tabs)

    def current(self):
        return self.TABS[self.tabs.currentIndex()][0]

    def load(self):
        for key, (t, maps) in self.tables.items():
            rows = self.win.db.list_rows(key)
            t.setRowCount(len(rows))
            for r, row in enumerate(rows):
                it = QTableWidgetItem(row[0])
                it.setData(Qt.UserRole, row[0])  # original name, to track renames
                t.setItem(r, 0, it)
                for c, m in enumerate(maps, start=1):
                    t.setItem(r, c, code_item(row[c], m))

    def add(self):
        t, maps = self.tables[self.current()]
        r = t.rowCount()
        t.insertRow(r)
        t.setItem(r, 0, QTableWidgetItem(""))
        for c, m in enumerate(maps, start=1):
            first = next(iter(m))
            t.setItem(r, c, code_item(first, m))
        t.editItem(t.item(r, 0))

    def remove(self):
        t, _ = self.tables[self.current()]
        for r in sorted({i.row() for i in t.selectedIndexes()}, reverse=True):
            t.removeRow(r)

    def move(self, d):
        t, maps = self.tables[self.current()]
        r = t.currentRow()
        if r < 0 or not 0 <= r + d < t.rowCount():
            return
        a = [t.takeItem(r, c) for c in range(t.columnCount())]
        b = [t.takeItem(r + d, c) for c in range(t.columnCount())]
        for c in range(t.columnCount()):
            t.setItem(r, c, b[c])
            t.setItem(r + d, c, a[c])
        t.selectRow(r + d)

    def save(self):
        key = self.current()
        t, maps = self.tables[key]
        rows, renames = [], {}
        for r in range(t.rowCount()):
            it = t.item(r, 0)
            name = (it.text() if it else "").strip()
            if not name:
                continue
            old = it.data(Qt.UserRole)
            if old and old != name:
                renames[old] = name
            rows.append((name,) + tuple((t.item(r, c).data(Qt.UserRole) or "") for c in range(1, len(maps) + 1)))
        self.win.db.save_list(key, rows, renames)
        self.load()
        self.win.data_changed()
        info(self, "تم الحفظ.")


# ================================================================== mapping page
class MappingPage(QWidget):
    COLS = ["رقم الحساب", "اسم الحساب", "المجموعة (المستوى 4)", "القائمة", "بند القائمة", "بند التدفق النقدي",
            "فئة دخل الفروع", "مجموعة المصروف", "ثابت / متغير", "الرصيد الختامي (آخر شهر)", "النوع حسب رقم الحساب"]
    FIELDS = [None, "name", "grp", None, "fs_item", "cf_item", "branch_cat", "exp_group", "behavior", None, None]

    def __init__(self, win):
        super().__init__()
        self.win = win
        self._loading = False
        lay = QVBoxLayout(self)
        h = QLabel("اربط كل حساب ببند القائمة وبند التدفق النقدي، وحسابات الدخل بفئة دخل الفروع ومجموعة المصروف. "
                   "الحسابات الجديدة في أي ميزان تُضاف هنا تلقائيًا بدون ربط. حدد عدة صفوف ثم «تطبيق على المحدد» لربطها مرة واحدة.")
        h.setWordWrap(True)
        h.setStyleSheet("padding:6px; background:#eef4fb; border-radius:6px;")
        lay.addWidget(h)
        bar = QHBoxLayout()
        self.search = QLineEdit()
        self.search.setPlaceholderText("بحث بالرقم أو الاسم أو البند...")
        self.search.textChanged.connect(self.apply_filter)
        self.filter = QComboBox()
        self.filter.addItems(["كل الحسابات", "غير مربوط ببند قائمة", "ميزانية بلا بند تدفق",
                              "دخل بلا فئة فروع", "مصروف تشغيلي بلا مجموعة", "حسابات الميزانية", "حسابات الدخل",
                              "ربطه لا يتفق مع رقم الحساب", "كل الحسابات التي بها مشكلة"])
        self.filter.currentIndexChanged.connect(self.apply_filter)
        bar.addWidget(self.search, 2)
        bar.addWidget(self.filter, 1)
        lay.addLayout(bar)
        bulk = QHBoxLayout()
        self.bulk_field = QComboBox()
        for i, f in enumerate(self.FIELDS):
            if f in ("fs_item", "cf_item", "branch_cat", "exp_group", "behavior", "grp"):
                self.bulk_field.addItem(self.COLS[i], i)
        self.bulk_field.currentIndexChanged.connect(self._fill_bulk_values)
        self.bulk_value = QComboBox()
        self.bulk_value.setEditable(True)
        b = QPushButton("تطبيق على المحدد")
        b.clicked.connect(self.apply_bulk)
        bulk.addWidget(QLabel("تعيين"))
        bulk.addWidget(self.bulk_field)
        bulk.addWidget(QLabel("="))
        bulk.addWidget(self.bulk_value, 2)
        bulk.addWidget(b)
        bulk.addStretch()
        lay.addLayout(bulk)

        self.table = QTableWidget(0, len(self.COLS))
        self.table.setHorizontalHeaderLabels(self.COLS)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.table.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.SelectedClicked |
                                   QAbstractItemView.EditKeyPressed)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.setSortingEnabled(False)
        freeze(self.table, 2)          # account number and name stay visible when scrolling sideways
        db_lists = {4: "fs_items", 5: "cf_items", 6: "branch_cats", 7: "exp_groups"}
        for col, table in db_lists.items():
            self.table.setItemDelegateForColumn(col, ListDelegate(self.table, lambda t=table: self.win.db.list_names(t)))
        self.table.setItemDelegateForColumn(8, ListDelegate(self.table, lambda: ["ثابت", "متغير"]))
        self.table.itemChanged.connect(self.on_changed)
        lay.addWidget(self.table, 1)
        bottom = QHBoxLayout()
        self.count = QLabel()
        bottom.addWidget(self.count)
        bottom.addStretch()
        for text, fn in (("✨ اقتراح ربط لغير المربوط", self.suggest_unmapped),
                         ("استيراد الربط من Excel", self.import_excel), ("تصدير الربط إلى Excel", self.export_excel),
                         ("حذف الحسابات المحددة", self.delete_selected)):
            btn = QPushButton(text)
            btn.clicked.connect(fn)
            bottom.addWidget(btn)
        lay.addLayout(bottom)
        self._fill_bulk_values()

    def _fill_bulk_values(self):
        col = self.bulk_field.currentData()
        self.bulk_value.clear()
        names = {4: "fs_items", 5: "cf_items", 6: "branch_cats", 7: "exp_groups"}
        if col in names and self.win.db:
            self.bulk_value.addItems([""] + self.win.db.list_names(names[col]))
        elif col == 8:
            self.bulk_value.addItems(["", "ثابت", "متغير"])

    def load(self):
        self._loading = True
        db = self.win.db
        accounts = sorted(db.accounts().values(), key=lambda a: importer.natural_key(a.code))
        last = db.periods()[-1] if db.periods() else None
        tb = db.tb(last) if last else {}
        sections = db.list_attr("fs_items", "section")
        self.table.setRowCount(len(accounts))
        for r, a in enumerate(accounts):
            self.table.setItem(r, 0, ro_item(a.code))
            for c, fld in enumerate(self.FIELDS):
                if fld:
                    self.table.setItem(r, c, QTableWidgetItem(getattr(a, fld) or ""))
            sec = sections.get(a.fs_item, "")
            self.table.setItem(r, 3, ro_item("ميزانية" if sec in defaults.BS_SECTIONS else "دخل" if sec else ""))
            v = (tb[a.code][5] - tb[a.code][6]) if a.code in tb else 0.0
            it = ro_item("-" if abs(v) < 0.005 else f"{abs(v):,.2f} {'مدين' if v > 0 else 'دائن'}")
            it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            self.table.setItem(r, 9, it)
            self.table.setItem(r, 10, ro_item(""))
        self.table.resizeColumnsToContents()
        self.table.setColumnWidth(1, 260)
        for c in (4, 5, 6, 7):
            self.table.setColumnWidth(c, 220)
        self._loading = False
        self.color_rows()
        self.apply_filter()
        self._fill_bulk_values()

    def row_account(self, r):
        vals = {"code": self.table.item(r, 0).text()}
        for c, fld in enumerate(self.FIELDS):
            if fld:
                it = self.table.item(r, c)
                vals[fld] = it.text().strip() if it else ""
        return Account(**{f: vals.get(f, "") for f in ACCOUNT_FIELDS})

    def color_rows(self):
        self._loading = True
        db = self.win.db
        sections = db.list_attr("fs_items", "section")
        kinds = db.list_attr("branch_cats", "kind")
        pmap = diagnose.prefix_map(db)
        n_unmapped = n_wrong = 0
        for r in range(self.table.rowCount()):
            a = self.row_account(r)
            sec = sections.get(a.fs_item, "")
            problems = set()
            exp = diagnose.expected_type(a.code, pmap)
            wrong = bool(sec) and not diagnose.type_matches(exp, sec)
            n_wrong += wrong
            t_item = self.table.item(r, 10)
            t_item.setText(diagnose.TYPE_NAMES[exp] + ("  ✗ لا يتفق مع الربط" if wrong else ""))
            t_item.setForeground(QColor("#b42318" if wrong else "#344054"))
            t_item.setData(Qt.UserRole, wrong)
            if not sec:
                problems.add(4)
            elif sec in defaults.BS_SECTIONS and not a.cf_item:
                problems.add(5)
            elif sec in defaults.IS_SECTIONS:
                if a.branch_cat not in kinds:
                    problems.add(6)
                elif kinds.get(a.branch_cat) == "opex" and not a.exp_group:
                    problems.add(7)
            if 4 in problems:
                n_unmapped += 1
            self.table.item(r, 3).setText("ميزانية" if sec in defaults.BS_SECTIONS else "دخل" if sec else "")
            for c in (4, 5, 6, 7):
                self.table.item(r, c).setBackground(QColor("#fee4e2") if c in problems else QColor(0, 0, 0, 0))
        self.count.setText(f"عدد الحسابات: {self.table.rowCount()}   |   غير مربوط ببند قائمة: {n_unmapped}"
                           f"   |   ربطه لا يتفق مع رقمه: {n_wrong}")
        self._loading = False

    def on_changed(self, item):
        if self._loading:
            return
        self.win.db.save_accounts([self.row_account(item.row())])
        self.color_rows()
        self.win.data_changed(reload_pages=False)

    def apply_filter(self):
        q = self.search.text().strip().lower()
        mode = self.filter.currentIndex()
        for r in range(self.table.rowCount()):
            text = " ".join((self.table.item(r, c).text() if self.table.item(r, c) else "") for c in range(9)).lower()
            show = not q or q in text
            if show and mode:
                bg = {1: 4, 2: 5, 3: 6, 4: 7}
                if mode in bg:
                    it = self.table.item(r, bg[mode])
                    show = it.background().color().alpha() > 0
                elif mode == 5:
                    show = self.table.item(r, 3).text() == "ميزانية"
                elif mode == 6:
                    show = self.table.item(r, 3).text() == "دخل"
                elif mode == 7:
                    show = bool(self.table.item(r, 10).data(Qt.UserRole))
                elif mode == 8:
                    show = bool(self.table.item(r, 10).data(Qt.UserRole)) or any(
                        self.table.item(r, c).background().color().alpha() > 0 for c in (4, 5, 6, 7))
            self.table.setRowHidden(r, not show)

    def selected_rows(self):
        return sorted({i.row() for i in self.table.selectedIndexes() if not self.table.isRowHidden(i.row())})

    def apply_bulk(self):
        col = self.bulk_field.currentData()
        val = self.bulk_value.currentText().strip()
        rows = self.selected_rows()
        if not rows:
            info(self, "حدد صفًا أو أكثر أولًا.")
            return
        self._loading = True
        for r in rows:
            self.table.item(r, col).setText(val)
        self._loading = False
        self.win.db.save_accounts([self.row_account(r) for r in rows])
        self.color_rows()
        self.apply_filter()
        self.win.data_changed(reload_pages=False)

    def suggest_unmapped(self):
        db = self.win.db
        accounts = db.accounts()
        reference = db.accounts()  # learn only from accounts mapped before this run
        sections = db.list_attr("fs_items", "section")
        pmap = diagnose.prefix_map(db)
        changed, lines = [], []
        for code, a in sorted(accounts.items()):
            if sections.get(a.fs_item):
                continue
            src, why = diagnose.suggest(reference, sections, pmap, code, a.name)
            if not src:
                continue
            for fld in ("fs_item", "cf_item", "branch_cat", "exp_group", "behavior", "grp"):
                if not getattr(a, fld):
                    setattr(a, fld, getattr(src, fld))
            changed.append(a)
            lines.append(f"{code} {a.name} ← {a.fs_item}  ({why})")
        if not changed:
            info(self, "لا توجد حسابات غير مربوطة يمكن اقتراح ربط لها.")
            return
        db.save_accounts(changed)
        self.load()
        self.filter.setCurrentIndex(0)
        self.win.data_changed(reload_pages=False)
        info(self, f"تم اقتراح ربط {len(changed)} حساب — راجعها قبل اعتماد التقارير:\n\n" + "\n".join(lines[:40]))

    def delete_selected(self):
        rows = self.selected_rows()
        if not rows:
            return
        if QMessageBox.question(self, "تأكيد", f"حذف {len(rows)} حساب من دليل الربط؟") != QMessageBox.Yes:
            return
        self.win.db.delete_accounts([self.table.item(r, 0).text() for r in rows])
        self.load()
        self.win.data_changed(reload_pages=False)

    def export_excel(self):
        path, _ = QFileDialog.getSaveFileName(self, "تصدير الربط", "ربط الحسابات.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.sheet_view.rightToLeft = True
        ws.append(self.COLS[:9])
        for r in range(self.table.rowCount()):
            ws.append([(self.table.item(r, c).text() if self.table.item(r, c) else "") for c in range(9)])
        wb.save(path)
        info(self, f"تم الحفظ في:\n{path}")

    def import_excel(self):
        path, _ = QFileDialog.getOpenFileName(self, "استيراد الربط", "", "Excel (*.xlsx *.xlsm)")
        if not path:
            return
        try:
            rows = importer.read_table(path)
            header = [str(v or "").strip() for v in rows[0]]
            idx = {name: header.index(name) for name in self.COLS[:9] if name in header}
            if "رقم الحساب" not in idx:
                raise ValueError("أول صف يجب أن يحتوي عناوين الأعمدة كما في ملف التصدير (رقم الحساب، اسم الحساب، ...).")
            accounts = []
            existing = self.win.db.accounts()
            for row in rows[1:]:
                code = importer.to_number(row[idx["رقم الحساب"]])
                code = str(int(code)) if code is not None and float(code).is_integer() else str(row[idx["رقم الحساب"]] or "").strip()
                if not code:
                    continue
                a = existing.get(code, Account(code, ""))
                for c, fld in enumerate(self.FIELDS):
                    if fld and self.COLS[c] in idx:
                        setattr(a, fld, str(row[idx[self.COLS[c]]] or "").strip())
                accounts.append(a)
            self.win.db.save_accounts(accounts)
        except Exception as e:  # noqa: BLE001
            error(self, str(e))
            return
        self.load()
        self.win.data_changed(reload_pages=False)
        info(self, f"تم استيراد ربط {len(accounts)} حساب.")


# ================================================================== data page
class DataPage(QWidget):
    def __init__(self, win):
        super().__init__()
        self.win = win
        lay = QVBoxLayout(self)
        h = QLabel("كل شهر: ارفع ميزان مراجعة الشركة للشهر (رقم الحساب، الاسم، والأعمدة الستة)، وحركة الفروع للشهر. "
                   "الشهور السابقة تبقى محفوظة، والتقارير الشهرية ومن بداية السنة تُبنى منها تلقائيًا.")
        h.setWordWrap(True)
        h.setStyleSheet("padding:6px; background:#eef4fb; border-radius:6px;")
        lay.addWidget(h)
        bar = QHBoxLayout()
        for text, fn in (("📥 رفع ميزان شهر", self.import_tb), ("📥 رفع حركة فروع لشهر", self.import_branches),
                         ("📥 أعمار العملاء لشهر", lambda: self.import_aging("customers")),
                         ("📥 أعمار الموردين لشهر", lambda: self.import_aging("suppliers")),
                         ("🗑 حذف بيانات الشهر المحدد", self.delete_period)):
            b = QPushButton(text)
            b.setMinimumHeight(34)
            b.clicked.connect(fn)
            bar.addWidget(b)
        bar.addStretch()
        lay.addLayout(bar)
        self.table = QTableWidget(0, 9)
        self.table.setHorizontalHeaderLabels(["الشهر", "عدد حسابات الميزان", "توازن الافتتاحي", "توازن الحركة",
                                              "توازن الختامي", "حسابات غير مربوطة", "الفروع المدخلة",
                                              "أعمار العملاء", "أعمار الموردين"])
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.horizontalHeader().setSectionResizeMode(6, QHeaderView.Stretch)
        self.table.verticalHeader().setVisible(False)
        lay.addWidget(self.table, 1)
        self.log = QPlainTextEdit()
        self.log.setReadOnly(True)
        self.log.setMaximumHeight(140)
        lay.addWidget(self.log)

    def load(self):
        db = self.win.db
        accounts = db.accounts()
        branch_info = {}
        for p, b, n in db.branch_periods():
            branch_info.setdefault(p, []).append(b)
        cust_p, supp_p = set(db.aging_periods("customers")), set(db.aging_periods("suppliers"))
        periods = sorted(set(db.periods()) | set(branch_info) | cust_p | supp_p, reverse=True)
        self.table.setRowCount(len(periods))
        for r, p in enumerate(periods):
            tb = db.tb(p)
            sums = [sum((v[i] or 0) for v in tb.values()) for i in range(1, 7)]
            self.table.setItem(r, 0, ro_item(calc.month_label(p)))
            self.table.item(r, 0).setData(Qt.UserRole, p)
            self.table.setItem(r, 1, ro_item(len(tb) if tb else "—"))
            for c, (d, cr) in enumerate(((0, 1), (2, 3), (4, 5)), start=2):
                diff = sums[d] - sums[cr]
                it = ro_item("—" if not tb else ("✓" if abs(diff) < 1 else f"فرق {diff:,.2f}"))
                it.setForeground(QColor("#067647" if abs(diff) < 1 else "#b42318"))
                self.table.setItem(r, c, it)
            un = sum(1 for code in tb if not (accounts.get(code) and accounts[code].fs_item))
            it = ro_item(un if tb else "—")
            if un:
                it.setForeground(QColor("#b42318"))
            self.table.setItem(r, 5, it)
            self.table.setItem(r, 6, ro_item("، ".join(branch_info.get(p, [])) or "—"))
            self.table.setItem(r, 7, ro_item(f"{len(db.aging('customers', p))} عميل" if p in cust_p else "—"))
            self.table.setItem(r, 8, ro_item(f"{len(db.aging('suppliers', p))} مورد" if p in supp_p else "—"))

    def _default_month(self):
        ps = self.win.db.periods()
        if ps:
            return calc.month_add(ps[-1], 1)
        import datetime
        d = datetime.date.today()
        return f"{d.year:04d}-{d.month:02d}"

    def _read(self, path):
        try:
            return importer.read_table(path)
        except Exception as e:  # noqa: BLE001
            error(self, f"تعذر قراءة الملف:\n{e}")
            return None

    def import_tb(self):
        dlg = MonthDialog(self, "ميزان أي شهر؟", self._default_month())
        if dlg.exec() != QDialog.Accepted:
            return
        period = dlg.period()
        start = QSettings().value("last_dir", os.path.expanduser("~"))
        path, _ = QFileDialog.getOpenFileName(self, "ملف ميزان المراجعة", start, "Excel / CSV (*.xlsx *.xlsm *.csv *.txt)")
        if not path:
            return
        QSettings().setValue("last_dir", os.path.dirname(path))
        rows = self._read(path)
        if rows is None:
            return
        first, mapping = importer.guess_mapping(rows)
        cm = ColumnMapDialog(self, rows, importer.ROLES, first, mapping,
                             "حدد عمود كل بيان. يكفي رقم الحساب والاسم والأرصدة الختامية، والأفضل الأعمدة الستة.")
        if cm.exec() != QDialog.Accepted:
            return
        m = cm.mapping()
        if "code" not in m or not ({"cd", "cc"} & set(m) or {"od", "oc"} & set(m) and {"md", "mc"} & set(m)):
            error(self, "لازم تحدد رقم الحساب والأرصدة الختامية (أو الافتتاحية + الحركة).")
            return
        accounts = importer.build_accounts(rows, cm.first_row(), m)
        parents = importer.mark_parent_accounts(accounts)
        if parents and QMessageBox.question(
                self, "حسابات رئيسية",
                f"يوجد {len(parents)} حساب رئيسي (رقمه بداية لأرقام حسابات أخرى) وأرصدته غالبًا إجمالي للفرعية.\n"
                "استبعادها من الميزان؟") == QMessageBox.Yes:
            accounts = [a for a in accounts if a.code not in parents]
        if period in self.win.db.periods() and QMessageBox.question(
                self, "تأكيد", f"يوجد ميزان محفوظ لشهر {calc.month_label(period)}. استبداله؟") != QMessageBox.Yes:
            return
        new = self.win.db.save_tb(period, [(a.code, a.name, a.od, a.oc, a.md, a.mc, a.cd, a.cc) for a in accounts])
        msg = [f"تم حفظ ميزان {calc.month_label(period)}: {len(accounts)} حساب."]
        hints = diagnose.tb_imbalance_hints(
            {a.code: (a.name, a.od, a.oc, a.md, a.mc, a.cd, a.cc) for a in accounts})
        for h_ in hints[:15]:
            prefix = {"error": "⛔", "warn": "⚠", "info": "ℹ"}[h_.severity]
            who = f" [{h_.code} {h_.name}]" if h_.code else ""
            msg.append(f"{prefix}{who} {h_.issue} {h_.fix}")
        if len(hints) > 15:
            msg.append(f"… و{len(hints) - 15} ملاحظة أخرى — افتح تقرير «تشخيص الخلل».")
        if new:
            msg.append(f"• {len(new)} حساب جديد أضيف لشاشة الربط ويحتاج ربط.")
        self.log.appendPlainText("\n".join(msg))
        self.win.data_changed()
        info(self, "\n".join(msg))

    def import_branches(self):
        db = self.win.db
        dlg = QDialog(self)
        dlg.setWindowTitle("رفع حركة فروع")
        f = QFormLayout(dlg)
        md = MonthDialog(self, "", self.win.current_period() or self._default_month())
        f.addRow("السنة:", md.year)
        f.addRow("الشهر:", md.month)
        r1 = QRadioButton("ملف لفرع واحد")
        r2 = QRadioButton("ملف واحد فيه كل الفروع (عمود للفرع)")
        r1.setChecked(True)
        grp = QButtonGroup(dlg)
        grp.addButton(r1)
        grp.addButton(r2)
        f.addRow(r1)
        branch = QComboBox()
        branch.setEditable(True)
        branch.addItems(db.branches())
        f.addRow("الفرع:", branch)
        f.addRow(r2)
        r2.toggled.connect(lambda on: branch.setEnabled(not on))
        note = QLabel("المبالغ: صافي حركة الشهر لكل حساب (مدين − دائن)، أو عمودين مدين ودائن للحركة، "
                      "أو الأعمدة الستة للميزان.")
        note.setWordWrap(True)
        f.addRow(note)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        f.addRow(bb)
        if dlg.exec() != QDialog.Accepted:
            return
        period = md.period()
        multi = r2.isChecked()
        bname = branch.currentText().strip()
        if not multi and not bname:
            error(self, "اختر الفرع.")
            return
        start = QSettings().value("last_dir", os.path.expanduser("~"))
        path, _ = QFileDialog.getOpenFileName(self, "ملف حركة الفروع", start, "Excel / CSV (*.xlsx *.xlsm *.csv *.txt)")
        if not path:
            return
        rows = self._read(path)
        if rows is None:
            return
        first, mapping = importer.guess_mapping(rows)
        roles = (["branch"] if multi else []) + importer.ROLES
        h_first, h_map = importer.guess_by_headers(rows, ("branch", "code", "name", "md") if multi else ("code", "name", "md"))
        if h_map and (not mapping or "branch" in h_map or len(h_map) > 2):
            first = h_first
            if not {"md", "mc", "cd", "cc"} & set(mapping):
                mapping = {}
            mapping.update(h_map)
        cm = ColumnMapDialog(self, rows, roles, first, mapping,
                             "حدد رقم الحساب وعمود المبلغ. لو الملف فيه عمود صافي واحد ضعه في «حركة مدين» "
                             "(السالب يُعتبر دائن).")
        if cm.exec() != QDialog.Accepted:
            return
        m = cm.mapping()
        if "code" not in m or (multi and "branch" not in m):
            error(self, "لازم تحدد رقم الحساب" + (" والفرع." if multi else "."))
            return
        if "name" not in m:
            m["name"] = m["code"]
        groups = {}
        if multi:
            bcol = m.pop("branch")
            for r in rows[cm.first_row():]:
                b = str(r[bcol] or "").strip()
                if b:
                    groups.setdefault(b, []).append(r)
        else:
            groups[bname] = rows[cm.first_row():]
        total_new, report = 0, []
        for b, brows in groups.items():
            accs = importer.build_accounts(brows, 0, m)
            data = [(a.code, a.name, a.mov_net) for a in accs if abs(a.mov_net) > 0.004]
            total_new += len(db.save_branch_mov(period, b, data))
            report.append(f"{b}: {len(data)} حساب")
        msg = f"تم حفظ حركة الفروع لشهر {calc.month_label(period)}:\n" + "\n".join(report)
        if total_new:
            msg += f"\n• {total_new} حساب جديد أضيف لشاشة الربط ويحتاج ربط."
        unknown = [b for b in groups if b not in [x for x, _ in db.branch_rows()]]
        if unknown:
            msg += "\n• فروع جديدة أضيفت للإعدادات: " + "، ".join(unknown)
        self.log.appendPlainText(msg)
        self.win.data_changed()
        info(self, msg)

    def import_aging(self, table):
        title = "أعمار العملاء" if table == "customers" else "أعمار الموردين"
        dlg = MonthDialog(self, f"{title} لأي شهر؟", self.win.current_period() or self._default_month())
        if dlg.exec() != QDialog.Accepted:
            return
        period = dlg.period()
        start = QSettings().value("last_dir", os.path.expanduser("~"))
        path, _ = QFileDialog.getOpenFileName(self, f"ملف {title}", start, "Excel / CSV (*.xlsx *.xlsm *.csv *.txt)")
        if not path:
            return
        rows = self._read(path)
        if rows is None:
            return
        roles = AGING_ROLES[table]
        first, mapping = guess_aging(rows, roles)
        help_text = ("حدد عمود كل بيان. المتأخر يُحسب تلقائيًا: العهدة أكثر من 30 يوم، وغيرها أكثر من 60 يوم."
                     if table == "customers" else
                     "حدد عمود كل بيان. المتجاوز يُحسب من فترة الائتمان، ويُحوّل بسعر الصرف من الإعدادات.")
        cm = ColumnMapDialog(self, rows, roles, first, mapping, help_text)
        if cm.exec() != QDialog.Accepted:
            return
        m = cm.mapping()
        if "no" not in m and "aname" not in m:
            error(self, "لازم تحدد عمود الرقم أو الاسم.")
            return
        out = []
        for r in rows[cm.first_row():]:
            rec = {}
            for role, col in m.items():
                v = r[col] if col < len(r) else None
                key = "name" if role == "aname" else role
                if role in ("b0", "b1", "b2", "b3", "b4", "guarantee", "credit"):
                    rec[key] = importer.to_number(v) or 0.0
                else:
                    if isinstance(v, float) and v.is_integer():
                        v = int(v)
                    rec[key] = "" if v is None else str(v).strip()
            if not rec.get("no") and not rec.get("name"):
                continue
            low = (rec.get("name") or "").replace("إ", "ا").replace("أ", "ا")
            if not rec.get("no") and ("الاجمالي" in low or "اجمالي" in low or "total" in low.lower()):
                continue
            for b in ("b0", "b1", "b2", "b3", "b4"):
                rec.setdefault(b, 0.0)
            out.append(rec)
        if period in self.win.db.aging_periods(table) and QMessageBox.question(
                self, "تأكيد", f"يوجد {title} محفوظة لشهر {calc.month_label(period)}. استبدالها؟") != QMessageBox.Yes:
            return
        self.win.db.save_aging(table, period, out)
        if table == "customers":
            s_ = aging.customers_summary(out)
            msg = (f"تم حفظ {title} لشهر {calc.month_label(period)}: {s_['count']} عميل، الرصيد {s_['total']:,.0f}، "
                   f"المتأخر {s_['overdue']:,.0f} ({s_['need_count']} عميل مطلوب منهم مطابقة).")
        else:
            s_ = aging.suppliers_summary(out, aging.rates(self.win.db), self.win.ctx.base_currency)
            msg = (f"تم حفظ {title} لشهر {calc.month_label(period)}: {s_['count']} مورد، الإجمالي {s_['total']:,.0f}، "
                   f"المتجاوز لفترة الائتمان {s_['overdue']:,.0f}.")
            if s_["no_rate"]:
                msg += f"\n⚠ {s_['no_rate']} مورد بعملة بدون سعر صرف — أضف السعر في الإعدادات."
        self.log.appendPlainText(msg)
        self.win.data_changed()
        info(self, msg)

    def delete_period(self):
        r = self.table.currentRow()
        if r < 0:
            return
        p = self.table.item(r, 0).data(Qt.UserRole)
        if QMessageBox.question(self, "تأكيد", f"حذف ميزان وحركة فروع شهر {calc.month_label(p)}؟") == QMessageBox.Yes:
            self.win.db.delete_period(p)
            self.win.data_changed()



# ================================================================== report view (screen)
class ReportView(QScrollArea):
    """Shows a report as real tables (scroll sideways for wide monthly reports)."""
    KIND_BG = {"sub": "#f2f5f9", "tot": "#e6edf6"}
    PARA_COLORS = {"good": "#067647", "warn": "#b54708", "bad": "#b42318"}

    def __init__(self):
        super().__init__()
        self.setWidgetResizable(True)
        self.setStyleSheet("QScrollArea { background: white; border: none; }")

    def show_message(self, text):
        lbl = QLabel(text)
        lbl.setWordWrap(True)
        lbl.setStyleSheet("color:#667085; padding: 20px; background: white;")
        self.setWidget(lbl)

    def show_report(self, rep, company):
        box = QWidget()
        box.setStyleSheet("background: white;")
        v = QVBoxLayout(box)
        for text, size, bold in ((company, 15, True), (rep.title, 13, True), (rep.subtitle, 10, False)):
            lbl = QLabel(text)
            lbl.setAlignment(Qt.AlignCenter)
            lbl.setStyleSheet(f"color:{NAVY if bold else '#5b6475'}; font-size:{pt(size)}; font-weight:{700 if bold else 400};")
            v.addWidget(lbl)
        for b in rep.blocks:
            if isinstance(b, R.Para):
                lbl = QLabel(b.text)
                lbl.setWordWrap(True)
                lbl.setStyleSheet(f"color:{self.PARA_COLORS.get(b.cls, '#5b6475')}; font-weight:{700 if b.cls in self.PARA_COLORS else 400};")
                v.addWidget(lbl)
                continue
            if b.title:
                lbl = QLabel(b.title)
                lbl.setStyleSheet(f"color:{NAVY}; font-size:{pt(11)}; font-weight:700; margin-top:8px;")
                v.addWidget(lbl)
            v.addWidget(self._table(b))
        v.addStretch()
        self.setWidget(box)

    def _table(self, b):
        t = QTableWidget(len(b.rows), len(b.columns))
        t.setHorizontalHeaderLabels([str(c) for c in b.columns])
        t.verticalHeader().setVisible(False)
        t.setEditTriggers(QAbstractItemView.NoEditTriggers)
        t.setSelectionMode(QAbstractItemView.ContiguousSelection)
        t.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        t.setWordWrap(False)
        t.verticalHeader().setDefaultSectionSize(int(30 * zoom.level()))
        t.setStyleSheet(f"QHeaderView::section {{ background:{NAVY}; color:white; font-weight:600; padding:4px; }}"
                        "QTableWidget { gridline-color: #e4e7ec; }")
        bold = QFont()
        bold.setBold(True)
        for r, row in enumerate(b.rows):
            label = QTableWidgetItem(row.label)
            t.setItem(r, 0, label)
            if row.kind == "hdr" and not row.values:
                label.setFont(bold)
                label.setForeground(QColor(NAVY))
                t.setSpan(r, 0, 1, len(b.columns))
                continue
            for i, val in enumerate(row.values):
                f = render.value_format(b, row, i)
                it = QTableWidgetItem(render.cell_text(val, f))
                if not isinstance(val, str) and f != "t":
                    it.setTextAlignment(Qt.AlignAbsolute | Qt.AlignRight | Qt.AlignVCenter)
                    if isinstance(val, (int, float)) and val < -0.5 and f == "n":
                        it.setForeground(QColor("#b42318"))
                st = row.styles[i] if row.styles and i < len(row.styles) else None
                if not st and isinstance(val, str) and val[:1] in ("▲", "▼"):
                    st = "warn"
                if st in self.PARA_COLORS:
                    it.setForeground(QColor(self.PARA_COLORS[st]))
                    it.setFont(bold)
                t.setItem(r, i + 1, it)
            for c in range(len(b.columns)):
                it = t.item(r, c) or QTableWidgetItem("")
                if t.item(r, c) is None:
                    t.setItem(r, c, it)
                if row.kind in ("sub", "tot", "hdr"):
                    it.setFont(bold)
                if row.kind in self.KIND_BG:
                    it.setBackground(QColor(self.KIND_BG[row.kind]))
                if row.kind == "note":
                    it.setForeground(QColor("#5b6475"))
        t.resizeColumnsToContents()
        t.setColumnWidth(0, max(t.columnWidth(0), 220 if len(b.columns) < 7 or not b.col_fmts else 110))
        long_text = False
        for c in range(1, len(b.columns)):
            is_text = bool(b.col_fmts and c - 1 < len(b.col_fmts) and b.col_fmts[c - 1] == "t")
            if is_text and t.columnWidth(c) > 260:
                long_text = True
                t.setColumnWidth(c, 320)
            else:
                t.setColumnWidth(c, max(min(t.columnWidth(c), 170 if not is_text else 260), 95))
        if long_text:
            t.setWordWrap(True)
            t.resizeRowsToContents()
        t.horizontalHeader().setDefaultAlignment(Qt.AlignCenter)
        t.horizontalHeader().setMinimumHeight(34)
        h = t.horizontalHeader().height() + sum(t.rowHeight(r) for r in range(t.rowCount())) + 22
        # a long table scrolls inside itself so its header row stays visible; the label column is frozen
        avail = max(320, self.viewport().height() - 70)
        if h > avail:
            t.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
            h = avail
        t.setFixedHeight(h)
        freeze(t, 1)
        return t

# ================================================================== reports page
class ReportsPage(QWidget):
    def __init__(self, win):
        super().__init__()
        self.win = win
        lay = QHBoxLayout(self)
        split = QSplitter()
        left = QWidget()
        lv = QVBoxLayout(left)
        lv.setContentsMargins(0, 0, 0, 0)
        self.list = QListWidget()
        self.list.setStyleSheet("QListWidget::item { padding: 5px; }")
        for title, group in (("— تقارير الفترة —", R.PERIOD_REPORTS), ("— التحليل الشهري —", R.MONTHLY_REPORTS)):
            it = QListWidgetItem(title)
            it.setFlags(Qt.NoItemFlags)
            it.setForeground(QColor(NAVY))
            f = it.font()
            f.setBold(True)
            it.setFont(f)
            self.list.addItem(it)
            for key, name, _ in group:
                i = QListWidgetItem(name)
                i.setData(Qt.UserRole, key)
                self.list.addItem(i)
        self.list.currentItemChanged.connect(lambda *_: self.show_report())
        lv.addWidget(self.list)
        self.branch_lbl = QLabel("الفرع:")
        self.branch = QComboBox()
        self.branch.currentIndexChanged.connect(lambda *_: self.show_report())
        self.view_lbl = QLabel("العرض:")
        self.view = QComboBox()
        self.view.addItem("الفترة", "period")
        self.view.addItem("من بداية السنة", "ytd")
        self.view.currentIndexChanged.connect(lambda *_: self.show_report())
        self.supp_lbl = QLabel("عرض الموردين:")
        self.supp = QComboBox()
        self.supp.addItem("تفصيلي حسب النوع (محلي / خارجي / خدمات)", "type")
        self.supp.addItem("إجمالي", "total")
        self.supp.currentIndexChanged.connect(self._supp_changed)
        for w in (self.branch_lbl, self.branch, self.view_lbl, self.view, self.supp_lbl, self.supp):
            lv.addWidget(w)
        split.addWidget(left)
        self.browser = ReportView()
        split.addWidget(self.browser)
        split.setStretchFactor(1, 4)
        split.setSizes([260, 1000])
        lay.addWidget(split)
        self.current_report = None

    def _supp_changed(self):
        self.win.db.set_setting("supp_view", self.supp.currentData())
        self.show_report()
        self.win.dashboard.load()

    def load(self):
        self.supp.blockSignals(True)
        self.supp.setCurrentIndex(max(0, self.supp.findData(self.win.db.setting("supp_view", "type") or "type")))
        self.supp.blockSignals(False)
        cur = self.branch.currentText()
        self.branch.blockSignals(True)
        self.branch.clear()
        self.branch.addItems(self.win.ctx.branches if self.win.ctx else [])
        i = self.branch.findText(cur)
        self.branch.setCurrentIndex(max(0, i))
        self.branch.blockSignals(False)
        if not self.list.currentItem():
            self.list.setCurrentRow(1)
        self.show_report()

    def key(self):
        it = self.list.currentItem()
        return it.data(Qt.UserRole) if it else None

    def show_report(self):
        key = self.key()
        if not key:
            return
        self.branch.setVisible(key in R.NEEDS_BRANCH)
        self.branch_lbl.setVisible(key in R.NEEDS_BRANCH)
        self.view.setVisible(key in R.NEEDS_VIEW)
        self.view_lbl.setVisible(key in R.NEEDS_VIEW)
        self.supp.setVisible(key == "suppliers")
        self.supp_lbl.setVisible(key == "suppliers")
        params = self.win.params()
        if not params:
            self.browser.show_message("لا توجد بيانات بعد — ارفع ميزان شهر من شاشة «البيانات الشهرية».")
            self.current_report = None
            return
        params.branch = self.branch.currentText()
        params.view = self.view.currentData()
        params.supp_view = self.supp.currentData() or "type"
        try:
            rep = R.build(key, self.win.ctx, params)
        except Exception:  # noqa: BLE001
            self.browser.show_message("حدث خطأ أثناء بناء التقرير:\n" + traceback.format_exc())
            self.current_report = None
            return
        self.current_report = rep
        self.browser.show_report(rep, self.win.ctx.company)


# ================================================================== start dialog
class StartDialog(QDialog):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("نظام التقارير المالية")
        self.resize(560, 420)
        self.choice = None
        lay = QVBoxLayout(self)
        t = QLabel("نظام التقارير المالية والإدارية")
        t.setStyleSheet(f"font-size:{pt(16)}; font-weight:700; color:{NAVY};")
        t.setAlignment(Qt.AlignCenter)
        lay.addWidget(t)
        self.recent = QListWidget()
        for p in QSettings().value("recent", []) or []:
            if os.path.exists(p):
                it = QListWidgetItem(f"📁 {os.path.basename(p)}   —   {os.path.dirname(p)}")
                it.setData(Qt.UserRole, p)
                self.recent.addItem(it)
        self.recent.itemDoubleClicked.connect(lambda it: self._done(("open", it.data(Qt.UserRole))))
        lay.addWidget(QLabel("الشركات الأخيرة (ضغطتين للفتح):"))
        lay.addWidget(self.recent, 1)
        for text, key in (("➕ شركة جديدة", "new"), ("📂 فتح ملف شركة", "open"),
                          ("📊 استيراد شركة من ملف Excel (النموذج القديم)", "migrate"),
                          ("⚡ تحليل سريع لميزان واحد", "quick")):
            b = QPushButton(text)
            b.setMinimumHeight(38)
            b.clicked.connect(lambda _=None, k=key: self._done((k, None)))
            lay.addWidget(b)

    def _done(self, choice):
        self.choice = choice
        self.accept()


# ================================================================== main window
class SystemWindow(QMainWindow):
    def __init__(self, path):
        super().__init__()
        self.db = CompanyDB(path)
        self.ctx = None
        self.resize(1360, 860)
        self._remember(path)

        tb = QToolBar()
        tb.setMovable(False)
        tb.setStyleSheet("QToolBar { spacing: 6px; padding: 4px; } QToolButton { padding: 5px 10px; font-weight:600; }")
        self.addToolBar(tb)
        self.period = QComboBox()
        self.period.setMinimumWidth(130)
        self.nmonths = QComboBox()
        for n, label in ((1, "شهر"), (3, "ربع (3 أشهر)"), (6, "نصف سنة"), (12, "سنة")):
            self.nmonths.addItem(label, n)
        self.comp = QComboBox()
        self.comp.setMinimumWidth(170)
        for w in (self.period, self.nmonths, self.comp):
            w.currentIndexChanged.connect(self.params_changed)
        tb.addWidget(QLabel(" الفترة المنتهية في: "))
        tb.addWidget(self.period)
        tb.addWidget(QLabel(" طول الفترة: "))
        tb.addWidget(self.nmonths)
        tb.addWidget(QLabel(" المقارنة مع: "))
        tb.addWidget(self.comp)
        tb.addSeparator()
        for text, fn in (("📊 Excel (التقرير)", lambda: self.export("xlsx", False)),
                         ("📄 PDF (التقرير)", lambda: self.export("pdf", False)),
                         ("📚 Excel (كل التقارير)", lambda: self.export("xlsx", True)),
                         ("📚 PDF (كل التقارير)", lambda: self.export("pdf", True))):
            a = QAction(text, self)
            a.triggered.connect(fn)
            tb.addAction(a)
        tb.addSeparator()
        draft = QAction("📑 مسودة القوائم المالية", self)
        draft.triggered.connect(self.open_draft)
        tb.addAction(draft)
        spacer = QWidget()
        spacer.setSizePolicy(spacer.sizePolicy().horizontalPolicy().Expanding, spacer.sizePolicy().verticalPolicy())
        tb.addWidget(spacer)
        for text, tip, fn, keys in (("🔍－", "تصغير (Ctrl + -)", zoom.zoom_out, "Ctrl+-"),
                                    ("100%", "الحجم العادي (Ctrl + 0)", zoom.reset, "Ctrl+0"),
                                    ("🔍＋", "تكبير (Ctrl + =) — أو Ctrl مع عجلة الماوس", zoom.zoom_in, "Ctrl+=")):
            a = QAction(text, self)
            a.setToolTip(tip)
            a.setShortcut(keys)
            a.triggered.connect(fn)
            tb.addAction(a)
            if text == "100%":
                self.zoom_action = a
        self.zoom_action.setText(f"{zoom.level() * 100:.0f}%")
        zoom.on_change(self._zoom_changed)
        up = QAction("🔄 فحص التحديثات", self)
        up.triggered.connect(lambda: self.check_updates(manual=True))
        tb.addAction(up)
        sw = QAction("🔁 تبديل الشركة", self)
        sw.triggered.connect(self.switch_company)
        tb.addAction(sw)
        self.updater = UpdateChecker(self)
        self.updater.result.connect(self._update_result)
        self.updater.failed.connect(self._update_failed)
        self._manual_check = False

        central = QWidget()
        h = QHBoxLayout(central)
        self.nav = QListWidget()
        self.nav.setFixedWidth(210)
        self.nav.setStyleSheet(
            f"QListWidget {{ background:{NAVY}; color:white; font-size:{pt(11)}; border:none; }}"
            "QListWidget::item { padding:13px 10px; } QListWidget::item:selected { background:#1d6fb8; }")
        self.stack = QStackedWidget()
        self.reports_page = ReportsPage(self)
        self.dashboard = DashboardPage(self)
        self.aging_page = AgingPage(self)
        self.pages = [("📊 لوحة التحكم", self.dashboard), ("📈 التقارير", self.reports_page),
                      ("👥 العملاء والموردين", self.aging_page),
                      ("🗓 البيانات الشهرية", DataPage(self)), ("🔗 ربط الحسابات", MappingPage(self)),
                      ("📋 البنود", ListsPage(self)), ("⚙ الإعدادات", SettingsPage(self))]
        for title, page in self.pages:
            self.nav.addItem(title)
            self.stack.addWidget(page)
        self.nav.currentRowChanged.connect(self.stack.setCurrentIndex)
        h.addWidget(self.nav)
        h.addWidget(self.stack, 1)
        self.setCentralWidget(central)
        self.data_changed()
        self.nav.setCurrentRow(0 if self.db.periods() else 3)
        self.check_updates(manual=False)

    def _zoom_changed(self):
        self.zoom_action.setText(f"{zoom.level() * 100:.0f}%")
        self.setFont(QApplication.instance().font())
        self.nav.setStyleSheet(
            f"QListWidget {{ background:{NAVY}; color:white; font-size:{pt(11)}; border:none; }}"
            "QListWidget::item { padding:13px 10px; } QListWidget::item:selected { background:#1d6fb8; }")
        self.reports_page.show_report()
        self.dashboard.load()
        self.aging_page.load()

    # -------- helpers
    def _remember(self, path):
        s = QSettings()
        recent = [p for p in (s.value("recent", []) or []) if p != path]
        s.setValue("recent", [path] + recent[:9])
        s.setValue("last_company", path)

    def current_period(self):
        return self.period.currentData()

    def params(self):
        end = self.current_period()
        if not end:
            return None
        return R.Params(end=end, n=self.nmonths.currentData() or 1, comp=self.comp.currentData() or "prev")

    def data_changed(self, reload_pages=True):
        self.ctx = calc.Ctx(self.db)
        version = f"نسخة {BUILD}" if BUILD else "نسخة تطوير"
        self.setWindowTitle(f"نظام التقارير المالية — {self.ctx.company}   ({os.path.basename(self.db.path)})   [{version}]")
        cur, comp = self.period.currentData(), self.comp.currentData()
        for w in (self.period, self.comp):
            w.blockSignals(True)
        self.period.clear()
        for p in reversed(self.ctx.periods):
            self.period.addItem(calc.month_label(p), p)
        if cur:
            self.period.setCurrentIndex(max(0, self.period.findData(cur)))
        self.comp.clear()
        self.comp.addItem("الفترة السابقة", "prev")
        self.comp.addItem("نفس الفترة من السنة السابقة", "ly")
        self.comp.addItem("بدون مقارنة", "none")
        for p in reversed(self.ctx.periods):
            self.comp.addItem(f"فترة تنتهي في {calc.month_label(p)}", p)
        self.comp.setCurrentIndex(max(0, self.comp.findData(comp)) if comp else 0)
        for w in (self.period, self.comp):
            w.blockSignals(False)
        if reload_pages:
            for _, page in self.pages:
                page.load()
        else:
            self.reports_page.load()
            self.dashboard.load()
            self.aging_page.load()

    def params_changed(self):
        self.reports_page.show_report()
        self.dashboard.load()
        self.aging_page.load()

    # -------- updates
    def check_updates(self, manual):
        self._manual_check = manual
        self.updater.check()

    def _update_result(self, newer, build, notes, url):
        if not newer:
            if self._manual_check:
                info(self, f"عندك أحدث نسخة (رقم {BUILD})." if BUILD else
                     f"هذه نسخة تطوير. آخر نسخة منشورة رقم {build}.")
            return
        box = QMessageBox(self)
        box.setWindowTitle("نسخة جديدة متاحة")
        box.setIcon(QMessageBox.Information)
        box.setText(f"نسخة جديدة رقم {build} متاحة (عندك رقم {BUILD}).")
        box.setInformativeText("الجديد:\n" + (notes[:900] or "—") +
                               "\n\nبعد التنزيل: اقفل البرنامج، وفك الملف، واستبدل فولدر البرنامج القديم. "
                               "ملفات الشركات (.fsdb) لا تتأثر.")
        download = box.addButton("تنزيل الآن", QMessageBox.AcceptRole)
        box.addButton("لاحقًا", QMessageBox.RejectRole)
        box.exec()
        if box.clickedButton() is download:
            QDesktopServices.openUrl(QUrl(url))

    def _update_failed(self, message):
        if self._manual_check:
            error(self, f"تعذر فحص التحديثات (تأكد من الاتصال بالإنترنت):\n{message}")

    def switch_company(self):
        QSettings().setValue("last_company", "")
        self.close()
        launch()

    def open_draft(self):
        if not self.current_period():
            info(self, "لا توجد بيانات بعد — ارفع ميزان شهر أولًا.")
            return
        from .draft_dialog import DraftDialog
        DraftDialog(self).exec()

    # -------- export
    def export(self, fmt, everything):
        params = self.params()
        if not params:
            info(self, "لا توجد بيانات بعد.")
            return
        rp = self.reports_page
        if self.stack.currentWidget() is self.dashboard and not everything:
            (self.dashboard.export_excel if fmt == "xlsx" else self.dashboard.export_pdf)()
            return
        if everything:
            reps = []
            for key, _, _ in R.ALL_REPORTS:
                p = R.Params(params.end, params.n, params.comp, rp.branch.currentText(), rp.view.currentData(),
                             rp.supp.currentData() or "type")
                try:
                    reps.append(R.build(key, self.ctx, p))
                except Exception:  # noqa: BLE001
                    traceback.print_exc()
            name = f"{self.ctx.company} - التقارير - {calc.month_label(params.end)}"
        else:
            if not rp.current_report:
                info(self, "اختر تقريرًا من شاشة التقارير.")
                return
            reps = [rp.current_report]
            name = f"{self.ctx.company} - {rp.current_report.title} - {calc.month_label(params.end)}"
        name = name.replace("/", "-").replace(":", " ")
        start = QSettings().value("last_dir", os.path.expanduser("~"))
        path, _ = QFileDialog.getSaveFileName(self, "حفظ", os.path.join(start, f"{name}.{fmt}"),
                                              "Excel (*.xlsx)" if fmt == "xlsx" else "PDF (*.pdf)")
        if not path:
            return
        try:
            if fmt == "xlsx":
                render.write_excel(reps, self.ctx.company, path)
            else:
                wide = any(max((len(b.columns) for b in r.blocks if hasattr(b, "columns")), default=0) > 8 for r in reps)
                write_pdf(render.reports_html(reps, self.ctx.company), path, landscape=wide)
        except Exception as e:  # noqa: BLE001
            error(self, str(e))
            return
        info(self, f"تم الحفظ في:\n{path}")


# ================================================================== launching
_windows = []


def _ask_new_path(parent, title, suggested):
    start = QSettings().value("last_dir", os.path.expanduser("~"))
    path, _ = QFileDialog.getSaveFileName(parent, title, os.path.join(start, suggested), "ملف شركة (*.fsdb)")
    if path and not path.lower().endswith(".fsdb"):
        path += ".fsdb"
    return path


def launch():
    last = QSettings().value("last_company", "")
    if last and os.path.exists(last):
        w = SystemWindow(last)
        w.show()
        _windows.append(w)
        return
    while True:
        dlg = StartDialog()
        if dlg.exec() != QDialog.Accepted or not dlg.choice:
            return
        kind, path = dlg.choice
        if kind == "quick":
            from .main_window import MainWindow
            w = MainWindow()
            w.show()
            _windows.append(w)
            return
        if kind == "open" and not path:
            start = QSettings().value("last_dir", os.path.expanduser("~"))
            path, _ = QFileDialog.getOpenFileName(None, "فتح ملف شركة", start, "ملف شركة (*.fsdb)")
            if not path:
                continue
        if kind == "new":
            path = _ask_new_path(None, "حفظ ملف الشركة الجديدة", "شركة جديدة.fsdb")
            if not path:
                continue
            if os.path.exists(path):
                os.remove(path)
        if kind == "migrate":
            src, _ = QFileDialog.getOpenFileName(None, "ملف Excel (النموذج القديم)", "", "Excel (*.xlsx *.xlsm)")
            if not src:
                continue
            try:
                company = migrate.inspect(src)["company"] or "شركة"
            except Exception as e:  # noqa: BLE001
                error(None, f"تعذر قراءة الملف:\n{e}")
                continue
            path = _ask_new_path(None, "حفظ ملف الشركة", f"{company}.fsdb")
            if not path:
                continue
            if os.path.exists(path):
                os.remove(path)
            QApplication.setOverrideCursor(Qt.WaitCursor)
            log = []
            try:
                db = CompanyDB(path)
                migrate.migrate(src, db, log=log.append)
                db.close()
            except Exception as e:  # noqa: BLE001
                QApplication.restoreOverrideCursor()
                error(None, f"فشل الاستيراد:\n{e}")
                continue
            QApplication.restoreOverrideCursor()
            info(None, "تم الاستيراد:\n" + "\n".join(log))
        w = SystemWindow(path)
        w.show()
        _windows.append(w)
        return
