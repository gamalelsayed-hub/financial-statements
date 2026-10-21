"""Draft financial statements on the synthetic company (no real data)."""
from fsapp.system import draft as DR
from fsapp.system.calc import Ctx
from tests.test_system import company  # noqa: F401  (fixture)


def test_draft_balances_ties_and_notes(company):  # noqa: F811
    ctx = Ctx(company)
    dr = DR.build(ctx, "2026-03")
    t = DR.bs_totals(dr.cur.bs)
    assert abs(t["diff"]) < 0.01 and abs(DR.bs_totals(dr.comp.bs)["diff"]) < 0.01
    assert dr.comp.end == "open:2026-01" and dr.comp.is_ is None      # no 2025 trial balance → opening balances
    assert abs(DR.is_totals(dr.cur.is_)["ni"] - dr.cur.ni) < 0.01
    assert abs(dr.cf_cur["diff"]) < 0.01 and dr.cf_cur["dep"] == 3000
    assert dr.cf_comp is None
    # default mapping
    m = dr.mapping
    assert m["النقد وما في حكمه"] == "cash" and m["العملاء وعهد المبيعات"] == "trade_rec"
    assert m["مجمع إهلاك الأصول الثابتة"] == "ppe" and m["قروض طويلة الأجل"] == "lt_fin"
    assert m["مصروفات تمويلية"] == "finance" and m["الإهلاك"] == "admin"
    # statements and notes: every note adds up to its line
    blocks, notes = DR.statements(dr, ctx)
    assert [k for k, _ in blocks] == ["fp", "is", "cf"]
    assert notes[0][0] == "is"                                        # income statement notes first
    for no, title, kind, rows, t1, t2 in DR.note_tables(dr, ctx, notes):
        assert abs(sum(r[2] for r in rows) - t1) < 0.01, title
        if t2 is not None:
            assert abs(sum(r[3] or 0 for r in rows) - t2) < 0.01, title
    assert DR.date_label("2026-06") == "30/6/2026" and DR.date_label("open:2026-01") == "31/12/2025"


def test_mapping_is_saved(company):  # noqa: F811
    ctx = Ctx(company)
    m = DR.load_map(ctx)
    m["المخزون"] = "other_rec"
    DR.save_map(company, m)
    ctx = Ctx(company)
    dr = DR.build(ctx, "2026-03")
    assert dr.mapping["المخزون"] == "other_rec" and dr.cur.bs["inventory"] == 0
    assert abs(DR.bs_totals(dr.cur.bs)["diff"]) < 0.01


def test_comparative_year_when_its_trial_balance_exists(company):  # noqa: F811
    company.set_setting("fiscal_start", "3")          # fiscal year Mar–Feb: Feb 2026 closes the previous year
    ctx = Ctx(company)
    dr = DR.build(ctx, "2026-03")
    assert dr.comp.end == "2026-02" and dr.comp.is_ is not None
    assert abs(DR.is_totals(dr.comp.is_)["ni"] - dr.comp.ni) < 0.01
    assert abs(dr.cf_cur["diff"]) < 0.01
    assert DR.date_label(dr.comp.end) == "28/2/2026"


def test_exports(company, tmp_path):  # noqa: F811
    import openpyxl
    from fsapp.system import draft_export as DX
    ctx = Ctx(company)
    dr = DR.build(ctx, "2026-03")
    opts = {"legal_form": "(شركة ذات مسؤولية محدودة)", "sign_title": "المدير المالي", "sign_name": "س"}
    wb = openpyxl.load_workbook(DX.to_excel(dr, ctx, str(tmp_path / "d.xlsx"), opts))
    assert wb.sheetnames[:4] == ["المركز المالي", "قائمة الدخل", "التدفقات النقدية", "الإيضاحات"]
    assert any(str(c.value).startswith("=SUM(") for row in wb["المركز المالي"].iter_rows() for c in row)
    page = DX.to_html(dr, ctx, opts)
    assert "قائمة المركز المالي كما في 31/3/2026" in page and "إيضاحات قائمة الدخل" in page
    assert page.index("إيضاحات قائمة الدخل") < page.index("إيضاحات قائمة المركز المالي")
