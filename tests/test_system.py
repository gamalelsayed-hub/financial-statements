"""Tests of the company system on a small synthetic company (no real data)."""
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fsapp.system import migrate, render  # noqa: E402
from fsapp.system import reports as R  # noqa: E402
from fsapp.system.calc import Ctx  # noqa: E402
from fsapp.system.db import Account, CompanyDB  # noqa: E402

# code, name, fs item, cf item, branch category, expense group, year opening (net), monthly movements Jan..Mar
CHART = [
    ("1101", "نقد بالصندوق", "النقد وما في حكمه", "النقد", "", "", 100000, None),
    ("1201", "العملاء", "العملاء وعهد المبيعات", "العملاء وعهد المبيعات", "", "", 50000, [20000, -5000, 10000]),
    ("1301", "المخزون", "المخزون", "المخزون", "", "", 80000, [-10000, 15000, -5000]),
    ("1401", "سيارات", "الأصول الثابتة (التكلفة)", "شراء واستبعاد أصول ثابتة", "", "", 60000, [0, 30000, -10000]),
    ("1402", "مجمع إهلاك سيارات", "مجمع إهلاك الأصول الثابتة", "الإهلاك", "", "", -20000, [-1000, -1000, 3000]),
    ("2101", "الموردون", "الموردون والدائنون التجاريون", "الموردون والدائنون التجاريون", "", "", -40000, [5000, -8000, 2000]),
    ("2201", "قرض طويل", "قروض طويلة الأجل", "القروض طويلة الأجل", "", "", -50000, [0, -20000, 5000]),
    ("3101", "رأس المال", "رأس المال", "رأس المال والاحتياطيات", "", "", -150000, [0, 0, 0]),
    ("3201", "أرباح محتجزة", "الأرباح المحتجزة", "تسويات حقوق الملكية", "", "", -30000, [0, 0, 0]),
    ("4101", "مبيعات", "إجمالي المبيعات", "", "المبيعات", "", 0, [-100000, -90000, -110000]),
    ("4102", "مردودات", "مردودات المبيعات", "", "مردودات المبيعات", "", 0, [2000, 1000, 3000]),
    ("4201", "إيراد بيع أصول", "إيرادات أخرى متعلقة بالنشاط", "أرباح بيع أصول ثابتة", "إيرادات أخرى", "", 0, [0, 0, -2000]),
    ("5101", "تكلفة المبيعات", "تكلفة المبيعات", "", "تكلفة المبيعات", "", 0, [70000, 63000, 77000]),
    ("5201", "رواتب إدارية", "رواتب وأجور - إدارية", "", "مصروفات تشغيلية", "الرواتب والأجور وما في حكمها", 0, [8000, 8000, 8500]),
    ("5202", "إيجار", "مصروفات عمومية وإدارية", "", "مصروفات تشغيلية", "الإيجارات", 0, [5000, 5000, 5000]),
    ("5203", "مصروف إهلاك", "الإهلاك", "", "مصروفات تشغيلية", "الإهلاك", 0, [1000, 1000, 1000]),
    ("5301", "فوائد", "مصروفات تمويلية", "", "مصروفات تمويلية", "", 0, [500, 500, 500]),
]
BRANCHES = {"الفرع أ": 0.6, "الفرع ب": 0.3}


def cash_movements():
    movs = [0.0, 0.0, 0.0]
    for row in CHART[1:]:
        for i in range(3):
            movs[i] -= row[7][i]
    return movs


def build_workbook(path):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "الإعدادات"
    s["B1"] = "الإعدادات والفهرس"
    s["C4"] = "شركة تجربة"
    import datetime
    s["C5"] = datetime.datetime(2026, 3, 31)
    s["C10"], s["C11"], s["C13"] = 0.2, 1000, 0.5
    s["B14"], s["E14"] = "الفرع", "العملة"
    for i, b in enumerate(BRANCHES, start=15):
        s.cell(i, 1, i - 14)
        s.cell(i, 2, b)
        s.cell(i, 3, "المنطقة")
    s["E15"], s["F15"] = "SAR", 1
    s = wb.create_sheet("البنود")
    from fsapp.system import defaults
    for i, (n, *_r) in enumerate(defaults.FS_ITEMS, start=5):
        s.cell(i, 1, n)
    for i, (n, *_r) in enumerate(defaults.CF_ITEMS, start=5):
        s.cell(i, 3, n)
    for i, (n, *_r) in enumerate(defaults.BRANCH_CATS, start=5):
        s.cell(i, 5, n)
    for i, n in enumerate(defaults.EXP_GROUPS, start=5):
        s.cell(i, 7, n)
    s = wb.create_sheet("الربط")
    s["A4"] = "رقم الحساب"
    for i, (code, name, fs, cf, cat, grp, *_r) in enumerate(CHART, start=5):
        for c, v in enumerate((int(code), name, "", "", fs, cf, cat, grp), start=1):
            s.cell(i, c, v or None)
    s = wb.create_sheet("المدخلات")
    s["A1016"] = "3) حركة الحسابات الشهرية — الشركة"
    s["A1018"] = "رقم الحساب"
    cash = cash_movements()
    for i, (code, name, *_r, opening, movs) in enumerate(CHART, start=1019):
        s.cell(i, 1, int(code))
        s.cell(i, 2, name)
        s.cell(i, 3, opening)
        for m, v in enumerate(movs if movs else cash):
            s.cell(i, 4 + m, v)
    s["A1421"] = "4) حركة الفروع الشهرية (الإيرادات والمصروفات)"
    s["A1423"] = "الفرع"
    r = 1424
    for b, share in BRANCHES.items():
        for code, name, fs, cf, cat, grp, opening, movs in CHART:
            if not cat or code == "5301" or code == "4201":
                continue
            s.cell(r, 1, b)
            s.cell(r, 2, int(code))
            s.cell(r, 3, name)
            for m, v in enumerate(movs):
                s.cell(r, 4 + m, v * share)
            r += 1
    wb.save(path)


@pytest.fixture()
def company(tmp_path):
    xlsx = tmp_path / "model.xlsx"
    build_workbook(xlsx)
    db = CompanyDB(str(tmp_path / "c.fsdb"))
    migrate.migrate(str(xlsx), db, log=lambda *_: None)
    return db


def test_migration_reads_everything(company):
    db = company
    assert db.setting("company") == "شركة تجربة"
    assert db.periods() == ["2026-01", "2026-02", "2026-03"]
    assert db.branches() == list(BRANCHES)
    assert len(db.accounts()) == len(CHART)
    assert len(db.list_names("fs_items")) == 51


def test_statements_balance_and_tie(company):
    ctx = Ctx(company)
    for months in (["2026-01"], ["2026-02"], ["2026-03"], ["2026-01", "2026-02", "2026-03"]):
        bs = ctx.bs_figures(months, 2)
        assert abs(bs["diff"]) < 0.01
        cf = ctx.cash_flow(months)
        assert abs(cf["cash_diff"]) < 0.01, cf
        assert not cf["no_item"]
    f = ctx.is_figures(["2026-01"])
    assert round(f["net_sales"]) == 98000
    assert round(f["gross"]) == 28000
    assert round(f["ni"]) == 28000 - 14000 - 500
    ytd = ctx.is_figures(["2026-03"], ytd=True)
    q = ctx.is_figures(["2026-01", "2026-02", "2026-03"])
    assert abs(ytd["ni"] - q["ni"]) < 0.01


def test_gain_on_sale_moves_to_investing(company):
    ctx = Ctx(company)
    cf = ctx.cash_flow(["2026-03"])
    op = {k: v for k, _, v in cf["sections"]["OP"]}
    inv = {k: v for k, _, v in cf["sections"]["INV"]}
    assert op["rc:أرباح بيع أصول ثابتة"] == -2000
    assert inv["rt:أرباح بيع أصول ثابتة"] == 2000


def test_branch_allocation(company):
    ctx = Ctx(company)
    cols, admin, comp = ctx.branches_report_data(["2026-01"])
    assert set(cols) == set(BRANCHES)
    # each branch got its share of the direct expenses; admin keeps the rest (10%)
    assert abs(cols["الفرع أ"]["net_sales"] - 98000 * 0.6) < 0.01
    assert abs(admin["direct"] - (-14000 * 0.1)) < 0.01
    assert abs(cols["الفرع أ"]["ratio"] - 2 / 3) < 1e-9
    total_after = sum(c["after"] for c in cols.values()) + admin["after"]
    assert abs(total_after - comp["after"]) < 0.01
    assert abs(comp["net"] - ctx.is_figures(["2026-01"])["ni"]) < 0.01


def test_all_reports_build_and_render(company, tmp_path):
    ctx = Ctx(company)
    reps = []
    for n, comp in ((1, "prev"), (3, "none"), (1, "ly")):
        p = R.Params(end="2026-03", n=n, comp=comp, branch="الفرع أ", view="ytd")
        for key, _, _ in R.ALL_REPORTS:
            reps.append(R.build(key, ctx, p))
    assert render.reports_html(reps[:20], ctx.company)
    render.write_excel(reps[:20], ctx.company, str(tmp_path / "out.xlsx"))
    checks = R.run_checks(ctx, R.Params(end="2026-03"))
    failed = [c[0] for c in checks if not c[3]]
    assert failed == [], failed


def test_new_accounts_from_a_trial_balance_are_unmapped(company):
    db = company
    new = db.save_tb("2026-04", [("9999", "حساب جديد", 0, 0, 100, 0, 100, 0)])
    assert new == ["9999"]
    assert db.accounts()["9999"].fs_item == ""
    ctx = Ctx(db)
    checks = {c[0]: c for c in R.run_checks(ctx, R.Params(end="2026-04"))}
    assert not checks["حسابات في الميزان غير مربوطة ببند قائمة"][3]


def test_list_rename_updates_accounts(company):
    db = company
    rows = [(n, s, r) for n, s, r in db.list_rows("fs_items")]
    rows = [("المخزون السلعي" if n == "المخزون" else n, s, r) for n, s, r in rows]
    db.save_list("fs_items", rows, {"المخزون": "المخزون السلعي"})
    assert db.accounts()["1301"].fs_item == "المخزون السلعي"
    db.save_accounts([Account("1301", "المخزون", fs_item="المخزون السلعي")])


# ---------------------------------------------------------------- diagnosis
from fsapp.system import diagnose as D  # noqa: E402


def _area_sum(findings, area):
    return sum(f.amount for f in findings if f.area == area and f.code)


def test_diagnosis_clean_company_has_no_errors(company):
    summary, findings = D.diagnose(Ctx(company), ["2026-03"])
    assert not [f for f in findings if f.severity == "error"], findings


def test_unmapped_account_explains_balance_sheet_and_cash_differences(company):
    db = company
    a = db.accounts()["1301"]
    a.fs_item, a.cf_item = "", ""
    db.save_accounts([a])
    ctx = Ctx(db)
    summary, findings = D.diagnose(ctx, ["2026-03"])
    bs = ctx.bs_figures(["2026-03"], 2)
    cf = ctx.cash_flow(["2026-03"])
    assert abs(bs["diff"]) > 1 and abs(cf["cash_diff"]) > 1
    assert abs(_area_sum(findings, "توازن المركز المالي") - bs["diff"]) < 0.01
    assert abs(sum(f.amount for f in findings if f.area == "التدفقات النقدية" and f.severity == "error")
               - cf["cash_diff"]) < 0.01
    # the unmapped account gets a suggestion from its name/number
    f = next(f for f in findings if f.area == "توازن المركز المالي")
    assert f.code == "1301" and "أصول" in f.issue


def test_balance_sheet_account_without_cash_flow_item(company):
    db = company
    a = db.accounts()["2201"]
    a.cf_item = ""
    db.save_accounts([a])
    ctx = Ctx(db)
    cf = ctx.cash_flow(["2026-02"])
    summary, findings = D.diagnose(ctx, ["2026-02"])
    errs = [f for f in findings if f.area == "التدفقات النقدية" and f.severity == "error"]
    assert [f.code for f in errs] == ["2201"]
    assert abs(sum(f.amount for f in errs) - cf["cash_diff"]) < 0.01


def test_revenue_number_mapped_to_balance_sheet_is_flagged(company):
    db = company
    a = db.accounts()["4102"]
    a.fs_item = "أرصدة دائنة أخرى"
    db.save_accounts([a])
    summary, findings = D.diagnose(Ctx(db), ["2026-03"])
    f = next(f for f in findings if f.code == "4102" and f.area == "تصنيف حسب رقم الحساب")
    assert f.severity == "error"


def test_prefix_scheme_is_configurable(company):
    db = company
    import json
    db.set_setting("prefix_map", json.dumps({"4": "X", "5": "R"}))
    pm = D.prefix_map(db)
    assert pm["4"] == "X" and pm["5"] == "R" and pm["1"] == "A"


def test_trial_balance_hints():
    rows = {
        "1101": ("نقد", 0, 0, 0, 0, 1000, 0),
        "2101": ("موردون", 0, 0, 0, 0, 300, 0),     # should be credit: posted on the wrong side
        "3101": ("رأس المال", 0, 0, 0, 0, 0, 700),
    }
    hints = D.tb_imbalance_hints(rows)
    assert any(h.code == "2101" and "نصف الفرق" in h.issue for h in hints)
    assert any(h.severity == "error" and "الختامي" in h.area for h in hints)


def test_suggest_from_closest_code(company):
    db = company
    db.ensure_accounts([("1202", "عميل جديد")])
    accounts = db.accounts()
    a, why = D.suggest(accounts, db.list_attr("fs_items", "section"), D.prefix_map(db), "1202", "عميل جديد")
    assert a.code == "1201" and a.fs_item == "العملاء وعهد المبيعات"


def test_keyword_suggestions_for_a_new_company(tmp_path):
    db = CompanyDB(str(tmp_path / "n.fsdb"))
    db.ensure_accounts([("1101", "الصندوق"), ("1210201", "مباني"), ("1210202", "مجمع اهلاك مباني"),
                        ("4101", "المبيعات"), ("5201", "رواتب الموظفين"), ("5202", "ايجار المعرض")])
    accounts = db.accounts()
    sections = db.list_attr("fs_items", "section")
    pm = D.prefix_map(db)

    def fs(code):
        a, _ = D.suggest(accounts, sections, pm, code, accounts[code].name)
        return a.fs_item if a else None

    assert fs("1101") == "النقد وما في حكمه"
    assert fs("1210202") == "مجمع إهلاك الأصول الثابتة"
    assert fs("4101") == "إجمالي المبيعات"
    a, _ = D.suggest(accounts, sections, pm, "5202", "ايجار المعرض")
    assert a.exp_group == "الإيجارات" and a.branch_cat == "مصروفات تشغيلية"
    # accumulated depreciation wins over the neighbouring "buildings" mapping
    b = accounts["1210201"]
    b.fs_item, b.cf_item = "الأصول الثابتة (التكلفة)", "شراء واستبعاد أصول ثابتة"
    db.save_accounts([b])
    accounts = db.accounts()
    a, _ = D.suggest(accounts, sections, pm, "1210202", "مجمع اهلاك مباني")
    assert a.fs_item == "مجمع إهلاك الأصول الثابتة"


# ---------------------------------------------------------------- aging and dashboard
from fsapp.system import aging as AG  # noqa: E402
from fsapp.system import dashboard  # noqa: E402


def test_overdue_rules():
    custody = {"ctype": "عهدة", "b0": 100, "b1": 50, "b2": 30, "b3": 0, "b4": 20}
    credit = dict(custody, ctype="آجل")
    assert AG.customer_overdue(custody) == 100   # over 30 days
    assert AG.customer_overdue(credit) == 50     # over 60 days
    base = {"b0": 10, "b1": 20, "b2": 30, "b3": 40, "b4": 50}
    assert AG.supplier_overdue(dict(base, credit=0)) == 140
    assert AG.supplier_overdue(dict(base, credit=60)) == 120
    assert AG.supplier_overdue(dict(base, credit=120)) == 50
    assert AG.supplier_overdue(dict(base, credit=180)) == 0


def test_aging_reports_and_checks(company):
    db = company
    db.save_aging("customers", "2026-03", [
        {"no": "1", "name": "عميل أ", "branch": "الفرع أ", "rep": "مندوب", "ctype": "آجل", "b0": 30000, "b1": 20000,
         "b2": 15000, "b3": 0, "b4": 0},
        {"no": "1", "name": "عميل مكرر الرقم", "branch": "الفرع ب", "rep": "مندوب", "ctype": "عهدة", "b0": 5000,
         "b1": 5000, "b2": 0, "b3": 0, "b4": 0}])
    db.save_currencies([("SAR", 1), ("USD", 3.75)])
    db.save_aging("suppliers", "2026-03", [
        {"no": "9", "name": "مورد محلي", "credit": 30, "currency": "SAR", "b0": 10000, "b1": 6000, "b2": 0, "b3": 0, "b4": 0},
        {"no": "10", "name": "مورد خارجي", "credit": 60, "currency": "USD", "b0": 1000, "b1": 0, "b2": 400, "b3": 0, "b4": 0}])
    ctx = Ctx(db)
    assert len(ctx.aging("customers", "2026-03")) == 2   # duplicate numbers are kept
    p = R.Params(end="2026-03")
    cust = R.build("customers", ctx, p)
    total_row = next(r for b in cust.blocks if isinstance(b, R.Table) for r in b.rows if r.label == "الإجمالي")
    assert total_row.values[1] == 75000 and total_row.values[3] == 15000 + 5000
    supp = R.build("suppliers", ctx, p)
    tot = next(r for b in supp.blocks if isinstance(b, R.Table) for r in b.rows if r.kind == "tot")
    assert tot.values[5] == 16000 + 1400 * 3.75 and tot.values[6] == 6000 + 400 * 3.75
    names = {c[0] for c in R.run_checks(ctx, p)}
    assert "تقرير أعمار العملاء = الميزان" in names and "تقرير أعمار الموردين = الميزان" in names


def test_dashboard_figures(company):
    ctx = Ctx(company)
    d = dashboard.compute(ctx, "2026-03")
    assert d["cards"]["profit"][0][1] == ctx.is_figures(["2026-03"])["net_sales"]
    cash = ctx.bs_figures(["2026-03"], 2)["cash"]
    assert abs(d["cards"]["cash"][0][1] - cash) < 0.01
    assert len(d["charts"]["months"]) == 12 and d["charts"]["profit"][0][1][3] is None
    moves = sum(v for _, v in d["charts"]["cash_moves"])
    start = ctx.bs_figures(["2026-01"], 0)["cash"]
    assert abs(start + moves - cash) < 0.01
    assert d["wc_rows"][0][2] == ctx.bs_figures(["2026-03"], 2)["receivables"]
    assert d["alerts"] and d["reading"]
